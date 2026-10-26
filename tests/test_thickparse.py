import io

import pytest

from server import thickparse as tp

PHRASE = ("Номинальные толщины стенки по поясам: с первого по второй – 16,0 мм; с третьего по четвертый – 15,0 мм; с пятого по седьмой – 13,0 мм; "
          "с восьмого по девятый – 12,0 мм; с десятого по двенадцатый – 11,0 мм.")
WANT = [16, 16, 15, 15, 13, 13, 13, 12, 12, 11, 11, 11]


def test_phrase_from_reference_report():
    r = tp.parse_document(PHRASE, n_belts=12)
    assert r["columns"][0]["values"] == WANT and r["columns"][0]["how"] == "текст"


def test_roman_table_like_reference_report():
    rows = [["Таблица 2 – Толщина элементов"], ["Обозначение элемента", "Номинальная толщина стенки, мм", "Расчетная толщина стенки, мм"],
            ["I пояс", "16,0", "16,0"], ["II пояс", "16,0", "15,5"], ["III пояс", "15,0", "15,0"], ["IV пояс", "15,0", "14,2"]]
    r = tp.parse_document(tables=[rows])
    names = [c["name"] for c in r["columns"]]
    assert names == ["Номинальная толщина стенки, мм", "Расчетная толщина стенки, мм"]
    assert r["columns"][0]["values"] == [16, 16, 15, 15] and r["columns"][1]["values"] == [16, 15.5, 15, 14.2]


def test_pasted_excel_table_with_min_max():
    text = "Пояс\tТолщина мин., мм\tТолщина макс., мм\n1\t14,8\t16,2\n2\t14,9\t16,0\n3\t13,8\t15,1\n"
    r = tp.parse_document(text, n_belts=3)
    assert [c["values"] for c in r["columns"]] == [[14.8, 14.9, 13.8], [16.2, 16.0, 15.1]]


def test_other_phrase_forms():
    assert tp.from_phrases("1-2 пояса – 16 мм, 3-4 пояса: 15 мм") == {1: 16, 2: 16, 3: 15, 4: 15}
    assert tp.from_phrases("пояс 1: 12 мм; пояс 2: 11,5 мм") == {1: 12, 2: 11.5}
    assert tp.from_phrases("с 1 по 3 - 10 мм") == {1: 10, 2: 10, 3: 10}


def test_xlsx_and_docx_files():
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    ws.append(["Пояс", "Толщина, мм"]); [ws.append([i, 16 - i // 3]) for i in range(1, 7)]
    b = io.BytesIO(); wb.save(b)
    text, tables = tp.read_file("УЗТ.xlsx", b.getvalue())
    assert tp.parse_document(text, tables, 6)["columns"][0]["values"] == [16, 16, 15, 15, 15, 14]
    import docx
    d = docx.Document(); d.add_paragraph(PHRASE)
    b = io.BytesIO(); d.save(b)
    text, tables = tp.read_file("паспорт.docx", b.getvalue())
    assert tp.parse_document(text, tables, 12)["columns"][0]["values"] == WANT


def test_not_found_and_mismatch_are_explained():
    with pytest.raises(ValueError, match="не найдены"):
        tp.parse_document("Резервуар вертикальный стальной, объём 10000 м3")
    with pytest.raises(ValueError, match="12 поясов, а в замерах 9"):
        tp.parse_document(PHRASE, n_belts=9)
    with pytest.raises(ValueError, match="не все пояса"):
        tp.parse_document("с первого по второй – 16,0 мм; с четвертого по пятый – 15,0 мм")
    with pytest.raises(ValueError):
        tp.read_file("scan.pdf", b"%PDF")


def test_implausible_values_ignored():
    assert tp.from_phrases("пояс 1: 1600 мм") == {}
