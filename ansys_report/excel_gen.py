"""Excel «Нивелировка»: облако точек стенки по результатам геодезической съёмки.

Структура листов повторяет образец «Нивелировка Т001-Т004.xlsx»:
  * матрица отклонений (мм) — вход (синий шрифт), пояса по столбцам;
  * блоки по поясам: угол, R новый, координаты Z/X/Y для импорта в Ansys (3d=true, Fit=false);
  * лист «Проверка» — допуск 1/200 высоты пояса.
Всё считается живыми формулами Excel.
"""
from __future__ import annotations
import math
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from .project import Project

BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MAT_COL0 = 13  # M — первый столбец матрицы отклонений


def _sheet(wb, name: str, p: Project, dev: list[list[float]]):
    t, s = p.tank, p.raw["survey"]
    nb, n = len(dev), s["points_per_belt"]
    ws = wb.create_sheet(name)
    ws["B1"] = f"={t['diameter']}/2"; ws["C1"] = "=B1*2"
    ws["F1"], ws["G1"], ws["H1"] = "Z", "X", "Y"
    ws["F2"] = "3d=true"
    ws["C3"] = t["belt_height"]; ws["C3"].font = BLUE
    ws["D3"], ws["E3"], ws["F3"] = "угол ", "R новый", "Fit=false"
    ws["K1"] = "шаг угла, °"; ws["L1"] = f"=360/{n}"
    ws["N1"] = "старт, °"; ws["O1"] = s["start_angle"]; ws["O1"].font = BLUE
    ws["L2"] = "точка / пояс →"
    # матрица отклонений
    for b in range(nb):
        c = ws.cell(3, MAT_COL0 + b, b + 1); c.font = BOLD; c.fill = HEAD; c.alignment = Alignment(horizontal="center")
    for i in range(n):
        ws.cell(4 + i, 12, i + 1)
        for b in range(nb):
            c = ws.cell(4 + i, MAT_COL0 + b, dev[b][i]); c.font = BLUE
    for b in range(nb):  # замыкающая точка
        ws.cell(4 + n, MAT_COL0 + b, f"={L(MAT_COL0+b)}4")
    ws.cell(4 + n, 12, 1)
    # блоки точек
    block = n + 2
    for k in range(1, nb + 1):
        mark = 3 + (k - 1) * block
        ws.cell(mark, 1, k).font = BOLD
        if k > 1:
            ws.cell(mark, 3, f"=$C$3*{k}")
        col = L(MAT_COL0 + k - 1)
        for i in range(n + 1):
            r = mark + 1 + i
            ws.cell(r, 2, (i % n) + 1)
            ws.cell(r, 3, f"={col}{4 + i}")
            ws.cell(r, 4, "=$O$1" if i == 0 else f"=D{r-1}-$L$1")
            ws.cell(r, 5, f"=$B$1+C{r}")
            ws.cell(r, 6, f"=ROUND(E{r}*SIN(RADIANS(D{r})),1)")
            ws.cell(r, 7, f"=$C${mark}")
            ws.cell(r, 8, f"=ROUND(E{r}*COS(RADIANS(D{r})),1)")
    for col, w in zip("ABCDEFGH", (6, 6, 9, 9, 10, 11, 9, 11)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    return ws, block


def _check_sheet(wb, p: Project, names: dict[str, list[list[float]]], n: int):
    t = p.tank
    ws = wb.create_sheet("Проверка")
    ws["A1"] = f"Отклонение стенки от вертикали (допуск 1/200 высоты) — {t['name']} №{t['tag']}"; ws["A1"].font = BOLD
    ws["A2"] = "Шаг пояса, мм"; ws["B2"] = t["belt_height"]; ws["B2"].font = BLUE
    ws["D2"] = "Отклонение пояса k = max |dk − d1| по точкам (d1 — 1-й пояс, база от окрайки днища)"
    hdr = ["Пояс", "Высота, мм", "Допуск 1/200, мм"]
    for nm in names:
        hdr += [f"Откл. от вертикали, {nm}, мм", f"Вывод {nm}"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(4, j, h); c.font = BOLD; c.fill = HEAD; c.border = BOX
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    nb = len(next(iter(names.values())))
    base = L(MAT_COL0)
    for k in range(1, nb + 1):
        r = 4 + k
        ws.cell(r, 1, k); ws.cell(r, 2, f"=$B$2*A{r}"); ws.cell(r, 3, f"=B{r}/200")
        for idx, nm in enumerate(names):
            col = L(MAT_COL0 + k - 1)
            cm, cv = 4 + 2 * idx, 5 + 2 * idx
            if k == 1:
                ws.cell(r, cm, "база"); ws.cell(r, cv, "—")
            else:
                rk = f"'{nm}'!{col}4:{col}{3 + n}"; r1 = f"'{nm}'!{base}4:{base}{3 + n}"
                ws.cell(r, cm, f"=SUMPRODUCT(MAX(ABS({rk}-{r1})))")
                ws.cell(r, cv, f'=IF({L(cm)}{r}<=C{r},"в допуске","ПРЕВЫШЕН")')
        for j in range(1, len(hdr) + 1):
            ws.cell(r, j).border = BOX
    for j in range(1, len(hdr) + 1):
        ws.column_dimensions[L(j)].width = 16
    ws.row_dimensions[4].height = 45


def build_excel(p: Project, out: Path) -> Path:
    wb = Workbook(); wb.remove(wb.active)
    tag = p.tank["tag"]
    n = p.raw["survey"]["points_per_belt"]
    states = {f"{tag} полный": p.survey_full, f"{tag} пустой": p.survey_empty}
    for name, dev in states.items():
        _sheet(wb, name, p, dev)
    _check_sheet(wb, p, states, n)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def export_points_txt(p: Project, dev: list[list[float]], out: Path) -> Path:
    """Текстовый файл точек (Z X Y, мм) для импорта в Ansys SpaceClaim/DesignModeler."""
    t, s = p.tank, p.raw["survey"]
    R0, n = t["diameter"] / 2, s["points_per_belt"]
    step = 360 / n
    lines = ["Z\tX\tY", "3d=true", "Fit=false"]
    for k, belt in enumerate(dev, 1):
        h = t["belt_height"] * k
        for i in range(n + 1):
            ang = math.radians(s["start_angle"] - step * i)
            R = R0 + belt[i % n]
            lines.append(f"{round(R*math.sin(ang),1)}\t{h}\t{round(R*math.cos(ang),1)}")
        lines.append("")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(lines), encoding="utf-8")
    return out
