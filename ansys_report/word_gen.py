"""Генерация технического отчёта (Word) по принятой форме технического отчёта."""
from __future__ import annotations
import re
from pathlib import Path
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from xml.sax.saxutils import escape as xml_escape
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from .calc import calc_all, fmt
from .project import Project
from . import text_ru as T

TEMPLATE = Path(__file__).parent / "templates" / "base.docx"
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII", "XIV", "XV"]
CENTER = WD_ALIGN_PARAGRAPH.CENTER
MATH_NS = ('xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
           'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"')


class Report:
    def __init__(self, p: Project):
        self.p, self.c = p, calc_all(p)
        self.d = docx.Document(TEMPLATE)
        self.fig = 0
        self.tab = 0
        self.missing: list[str] = []

    # ---------- примитивы (оформление как в образце ТО-019-26: TNR 12, по ширине, отступ 1 см) ----------
    FONT = "Times New Roman"

    @classmethod
    def _run(cls, par, text, size=12, bold=False, italic=False):
        r = par.add_run(text)
        r.font.name = cls.FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), cls.FONT)
        r.font.size = Pt(size); r.bold = bold; r.italic = italic
        r.font.color.rgb = RGBColor(0, 0, 0)
        return r

    def para(self, text="", style="Normal", bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, italic=False,
             size=12, indent=1.0, keep_next=False, space_before=0, space_after=0):
        par = self.d.add_paragraph(style=style)
        text = re.sub(r"(?i)(резервуар\w*) Резервуар\b", r"\1", text)      # название по умолчанию «Резервуар» не дублируем
        text = re.sub(r" {2,}", " ", text)
        if text:
            self._run(par, text, size, bold, italic)
        pf = par.paragraph_format
        par.alignment = align
        pf.first_line_indent = Cm(indent) if indent else None
        pf.space_before, pf.space_after = Pt(space_before), Pt(space_after)
        pf.line_spacing = 1.0
        pf.keep_with_next = keep_next
        return par

    def h1(self, text, new_page=False):
        """Заголовок раздела: полужирный 12 пт, абзацный отступ, отбивка 12/6 пт."""
        h = self.d.add_paragraph(style="Heading 1")
        self._run(h, text, 12, bold=True)
        h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h.paragraph_format.first_line_indent = Cm(1.0)
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(12), Pt(6)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.page_break_before = new_page
        return h

    def h2(self, text):
        """Подраздел (1.1, 2.1 …): полужирный 12 пт по ширине."""
        return self.para(text, bold=True, keep_next=True, space_before=6, space_after=3)

    def body(self, text): return self.para(text)

    def bullets(self, items):
        for it in items:
            self.para("- " + it, indent=1.25)

    def caption(self, text):
        return self.para(text, align=CENTER, indent=0, space_before=3, space_after=9)

    def page_break(self): self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def table(self, header: list[str], rows: list[list], widths=None, caption: str | None = None,
              fills: dict | None = None, size: int = 12):
        if caption:
            self.tab += 1
            self.para(f"Таблица {self.tab} – {caption}", align=WD_ALIGN_PARAGRAPH.LEFT, indent=0,
                      keep_next=True, space_before=6, space_after=3)
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            self._cell(t.rows[0].cells[i], h, bold=True, align=CENTER, size=size)
        self._repeat_header(t.rows[0])
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                self._cell(cells[i], str(v), align=CENTER if i else WD_ALIGN_PARAGRAPH.LEFT,
                           red=(str(v) == "Не выполнен"), size=size)
        if widths:
            t.autofit = False
            for i, w in enumerate(widths): t.columns[i].width = Cm(w)
            for r in t.rows:
                for i, w in enumerate(widths): r.cells[i].width = Cm(w)
        self.para(indent=0)
        return t

    @staticmethod
    def _cell(cell, text, bold=False, align=None, red=False, size=12):
        cell.text = ""
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(0); par.paragraph_format.first_line_indent = Cm(0)
        if align is not None: par.alignment = align
        r = Report._run(par, text, size, bold)
        if red: r.font.color.rgb = RGBColor(0xC0, 0, 0)

    @staticmethod
    def _repeat_header(row):
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)

    def figure(self, key: str, caption: str, width_cm=12.5, optional=False):
        """Вставляет картинку из images_dir (файл <key>.png|jpg); иначе — заметная заглушка (optional: рисунок пропускается)."""
        img = self._find_image(key)
        if optional and not img:
            return
        self.fig += 1
        if img:
            pp = self.para(indent=0, align=CENTER, keep_next=True, space_before=6)
            pp.add_run().add_picture(str(img), width=Cm(width_cm))
        else:
            self.missing.append(key)
            par = self.para(f"[Рисунок не найден: {key}.png — положите экспорт из Ansys в папку images]",
                            align=CENTER, italic=True, indent=0)
            par.runs[0].font.color.rgb = RGBColor(0xC0, 0, 0)
        self.caption(f"Рисунок {self.fig}. {caption}")

    def _find_image(self, key):
        folder = self.p.root / self.p.raw.get("images_dir", "images")
        mapped = (self.p.raw.get("images") or {}).get(key)   # явное соответствие рисунок -> файл из Ansys
        if mapped and (folder / mapped).exists():
            return folder / mapped
        for ext in (".png", ".jpg", ".jpeg"):
            f = folder / f"{key}{ext}"
            if f.exists(): return f
        return None

    def math(self, text):
        return self.para(text, align=CENTER, italic=True, indent=0, space_before=3, space_after=3)

    # ---------- части отчёта ----------
    def header_footer(self):
        r, t = self.p.report, self.p.tank
        sec = self.d.sections[0]
        tbl = sec.header.tables[0]
        self._set_cell_text(tbl.rows[0].cells[0].paragraphs[0], t["name"])
        self._set_cell_text(tbl.rows[0].cells[2].paragraphs[1], f"№ {r['number']}")
        fp = sec.footer.paragraphs[0]
        runs = [x for x in fp.runs if x.text]
        runs[0].text = str(r["year"]); runs[1].text = ""  # «202»+«6»

    @staticmethod
    def _set_cell_text(par, text):
        runs = [x for x in par.runs if x.text is not None]
        runs[0].text = text
        for x in runs[1:]: x.text = ""

    def title_page(self):
        """Титульный лист как в образце: всё 14 пт, по центру; заголовок полужирный."""
        r, t = self.p.report, self.p.tank
        C = dict(align=CENTER, indent=0, size=14)
        for _ in range(6): self.para(**C)
        self.para(f"ТЕХНИЧЕСКИЙ ОТЧЕТ №{r['number']}", bold=True, **C)
        self.para("по результатам расчета методом конечных элементов", **C)
        self.para(**C)
        self.para(r["title"], bold=True, **C)
        for _ in range(9): self.para(**C)
        tb = self.d.add_table(rows=1, cols=2)
        left, right = tb.rows[0].cells
        left.text = ""; right.text = ""
        self._run(left.paragraphs[0], "Выполнил:", 14)
        name = (r.get("executor_name") or "").strip()
        lines = [(r.get("executor_position") or "").strip(), f"__________ / {name or '____________'} /",
                 f"«___» ___________ {r['year']} г."]
        right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self._run(right.paragraphs[0], lines[0], 14)
        for ln in lines[1:]:
            pp = right.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            self._run(pp, ln, 14)
        # год — в рамке, прижатой к нижнему краю страницы (не зависит от числа строк выше)
        yp = self.para(f"{r['year']} г.", **C)
        fr = OxmlElement("w:framePr")
        for k, v in (("w:w", "9000"), ("w:hSpace", "0"), ("w:wrap", "around"), ("w:vAnchor", "margin"),
                     ("w:hAnchor", "margin"), ("w:xAlign", "center"), ("w:yAlign", "bottom")):
            fr.set(qn(k), v)
        yp._p.get_or_add_pPr().insert(0, fr)
        self.page_break()

    def toc(self):
        self.para("СОДЕРЖАНИЕ", bold=True, align=CENTER, indent=0, size=14, space_after=6)
        par = self.d.add_paragraph()
        run = par.add_run()
        for typ, txt in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None)):
            if typ:
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), typ); run._r.append(e)
            else:
                e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; run._r.append(e)
        run.add_text("Правый клик → «Обновить поле», чтобы построить оглавление.")
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); run._r.append(e)
        self.page_break()
        # просим Word обновить поля при открытии
        st = self.d.settings.element
        uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); st.append(uf)

    def sec1(self):
        r = self.p.report
        self.h1("1 Вводная часть")
        self.para("1.1. Прочностной расчет проведен в соответствии с нормативными документами:", bold=True)
        self.bullets(T.NORMS_SHORT)
        if (r.get("contractor") or "").strip() or r.get("contractor_phone"):     # без сведений раздел не выводим
            self.para("1.2 Сведения о подрядчике, проводившем прочностной расчет", bold=True)
            if (r.get("contractor") or "").strip():
                self.body(r["contractor"].strip())
            if r.get("contractor_phone"): self.body(f"Тел. {r['contractor_phone']}.")

    # ---------- формулы Word (OMML) и мелочи оформления ----------
    @staticmethod
    def _m(parts) -> str:
        """Части формулы: строка - текст, кортеж (числитель, знаменатель) - дробь (части рекурсивно)."""
        if isinstance(parts, str):
            parts = [parts]
        out = []
        for x in parts:
            if isinstance(x, tuple):
                out.append(f"<m:f><m:num>{Report._m(x[0])}</m:num><m:den>{Report._m(x[1])}</m:den></m:f>")
            else:
                out.append('<m:r><w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/><w:sz w:val="24"/></w:rPr>'
                           f'<m:t xml:space="preserve">{xml_escape(str(x))}</m:t></m:r>')
        return "".join(out)

    def formula(self, parts):
        """Формула редактора формул Word по центру (как в эталоне: R = 325∙0.7∙1.0 / 1.05∙1.05 = 206,35 МПа)."""
        par = self.para(indent=0, align=CENTER, space_before=3, space_after=3)
        par._p.append(parse_xml(f'<m:oMathPara {MATH_NS}><m:oMath>{self._m(parts)}</m:oMath></m:oMathPara>'))
        return par

    @staticmethod
    def _cell_rich(cell, chunks, size=12, align=None):
        """Ячейка с верхними индексами: chunks = [(текст, верхний_индекс)]."""
        cell.text = ""
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(0); par.paragraph_format.first_line_indent = Cm(0)
        if align is not None: par.alignment = align
        for text, sup in chunks:
            r = Report._run(par, text, size)
            r.font.superscript = bool(sup)

    @staticmethod
    def _merge_text(a, b, text, bold=True, size=12):
        c = a.merge(b)
        Report._cell(c, text, bold=bold, align=CENTER, size=size)
        return c

    @staticmethod
    def _pow10(v: float) -> list:
        """2,06·10^11 как в эталоне: [('2,06х10', False), ('11', True)]."""
        import math
        e = int(math.floor(math.log10(abs(v)))) if v else 0
        mant = v / 10 ** e
        return [(f"{mant:.2f}".rstrip("0").rstrip(",.").replace(".", ",") + "х10", False), (str(e), True)]

    @property
    def rd(self) -> dict:
        return self.p.raw.get("report_data") or {}

    # ---------- 2 Информация об объекте ----------
    def sec2(self):
        t, m, c, rd = self.p.tank, self.p.material, self.c, self.rd
        th = t["belts_thickness"]
        n = len(th)
        self.h1("2 Информация об объекте моделирования")
        where = f" установлен на объекте: «{t['site']}»" + (f" (№ {t['site_number']})" if t.get("site_number") else "") if t.get("site") else ""
        self.body(f"Резервуар {t['name']}{where}" + ((f" и предназначен для {t['purpose']}." if where else f" предназначен для {t['purpose']}.") if t.get("purpose") else "."))
        self.body("Габариты резервуара:")
        self.bullets([f"диаметр {t['diameter']} мм;", f"высота стенки {t['wall_height']} мм."])
        self.body("Общий вид резервуара показан на рисунке 1.")
        if t.get("fill_level"):
            rho_t = float(t.get("product_density") or 0) / 1000
            self.body(f"Верхний аварийный уровень налива продукта по проекту составляет – {fmt(t['fill_level']/1000)} м. "
                      f"Плотность хранимого продукта составляет {fmt(rho_t, 1 if abs(rho_t * 10 - round(rho_t * 10)) < 1e-9 else 3)} т/м3.")
        groups = T.thickness_groups(th).replace("первого –", "первого пояса стенки резервуара –", 1)
        steel = f"Стенка резервуара состоит из листов стали {m['name']}" + (f", {m['standard']}" if (m.get("standard") or "").strip() else "") + ". "
        self.body(steel + "Номинальная толщина " + groups + ".")
        self.body("Предельное отклонение от вертикали каждого листа любого пояса стенки ограничивается величиной 1/200 вертикального расстояния "
                  "от окрайки днища до контролируемого пояса стенки. Соответствующие значения предельных отклонений для "
                  f"I - {ROMAN[n-1]} поясов стенки составляют " + " – ".join(g3(x) for x in c["limits_dev"]) + " мм.")
        self.body("Отклонения стенки от вертикали в каждом поясе замерялись в плоскостях x = " + " – ".join(fmt(h, 0) for h in c["heights"]) + " мм.")
        self.body("Основные характеристики РВС представлены в таблице 1. Толщины стенок по поясам представлены в таблице 2"
                  + (f" согласно {t['kmd']}." if (t.get("kmd") or "").strip() else "."))
        self.figure("fig01_general", "Общий вид объекта исследования", width_cm=11.3)
        rows = self.object_rows_ref() if rd else self.object_rows()
        self.table(["Наименование величин", "Значение"], rows, widths=[11, 5], caption="Характеристика объекта исследования")
        self.page_break()
        rows = [[f"{ROMAN[i]} пояс", fmt(x), fmt(x)] for i, x in enumerate(th)]
        if t.get("bottom_thickness"):
            rows.append(["Днище резервуара", fmt(t["bottom_thickness"]), fmt(t["bottom_thickness"])])
        tb = self.table(["Обозначение элемента", "Номинальная толщина стенки, мм", "Расчетная толщина стенки, мм"], rows,
                        caption="Толщина элементов", widths=[6, 5, 5])
        for r in tb.rows[1:]:
            r.cells[0].paragraphs[0].alignment = CENTER
        self.sec2_1()
        self.para("2.2 Необходимый состав расчета", bold=True, space_before=6)
        self.composition()

    def sec2_1(self):
        m, c, t = self.p.material, self.c, self.p.tank
        n = len(t["belts_thickness"])
        sy = m["yield_mpa"]
        self.para("2.1 Механические характеристики материалов, используемых в расчете", bold=True, space_before=6)
        std = f" согласно {m['standard']}" if (m.get("standard") or "").strip() else ""
        self.body(f"Материал резервуара – сталь {m['name']}. Характеристики материала{std} приведены в таблице {self.tab + 1}.")
        tb = self.table(["Сталь", "Плотность, кг/м3", "Модуль Юнга, Па", "Коэффициент Пуассона", "Предел прочности, МПа",
                         "Предел текучести, МПа", "Для толщин стенок, мм"],
                        [[m["name"], fmt(m["density"], 0), "", g3(m["poisson"]), fmt(m["ultimate_mpa"], 0), fmt(sy, 0), m.get("thickness_range") or "–"]],
                        caption="механические характеристики используемых материалов", widths=[2.4, 2.4, 2.6, 2.7, 2.4, 2.4, 2.4], size=11)
        self._cell_rich(tb.rows[1].cells[2], self._pow10(float(m["young_modulus_mpa"]) * 1e6), size=11, align=CENTER)
        self.body("Согласно ГОСТ 31385-2023 и СП 16.13330.2020:")
        for label, gc, key in ((f"Расчетные сопротивления для 1го пояса:", "0.7", "R1"), (f"Расчетные сопротивления для 2-{n} го поясов:", "0.8", "R2")):
            self.body(label)
            self.formula(["R=", (f"{g3(sy)}∙{gc}∙1.0", "1.05∙1.05"), f"={fmt(c[key], 2)} МПа"])
        self.body("Зона сопряжения стенки с днищем (уторный узел) при необходимости расчета с учетом развития пластических деформаций:")
        self.formula(["R=", (f"{g3(sy)}∙1.2∙1.0", "1.05∙1.05"), f"={fmt(c['Rutor'], 2)} МПа"])
        last = T.ORD_NOM[n - 1] if n <= len(T.ORD_NOM) else f"{n}-й"
        self.body(f"Напряжения в первом поясе не должны превышать {int(c['R1'])} МПа, со второго по {last} пояс не должны превышать {int(c['R2'])} МПа. "
                  f"При субмоделировании уторного шва РВС напряжения не должны превышать {int(c['Rutor'])} МПа.")
        self.body(f"Расчеты устойчивости стенки резервуара {t['name']} с учетом отклонений от проектной формы выполнены с использованием "
                  "программно-вычислительного комплекса, реализующего метод конечных элементов.")
        if m.get("bilinear"):
            self.body(f"В модели материала стенки и окрайки днища резервуара использовалась упругопластическая модель с билинейной диаграммой "
                      f"деформирования. При этом предел текучести был принят равным σу = {g3(sy)} МПа, а модуль упругости второго участка диаграммы определен как:")
        else:
            self.body("Расчет НДС выполнен в физически линейной постановке (п. 10.2.2 ГОСТ Р 58622-2019): модель материала стенки, днища и крыши – "
                      f"линейно-упругая. Для расчета уторного узла с учетом развития пластических деформаций используется билинейная диаграмма "
                      f"деформирования с пределом текучести σу = {g3(sy)} МПа и модулем упругости второго участка:")
        e_mpa = float(m["young_modulus_mpa"])
        self.formula(["E2=", ("σвр-σу", [("∆", "100"), "-", ("σу", "E")]), f"={fmt(c['E2'], 1)} МПа"])
        self.body("где")
        for line in (f"σвр = {g3(m['ultimate_mpa'])} МПа – временное сопротивление стали {m['name']};", f"σу = {g3(sy)} МПа – предел текучести;",
                     f"Δ = {g3(m['elongation_pct'])} – относительное удлинение, %;"):
            self.body(line)
        par = self.body("Е = ")
        for text, sup in self._pow10(e_mpa * 1e6):
            self._run(par, text.replace("х", "·"), 12).font.superscript = sup
        self._run(par, " Па – модуль упругости.", 12)

    def composition(self):
        gost = Path(__file__).parent / "templates" / "gost"
        for x in T.COMPOSITION_REF:
            if isinstance(x, str):
                self.body(x)
            elif x[0] == "img":
                f = gost / x[1]
                if f.exists():
                    self.para(indent=0, align=CENTER, keep_next=True, space_before=6).add_run().add_picture(str(f), width=Cm(x[2]))
            elif x[0] == "cap":
                self.para(x[1], indent=0, align=CENTER, space_after=6)
            elif x[0] == "formula":
                self.para(x[1], indent=0, align=CENTER, space_before=3, space_after=3)
            elif x[0] == "item":
                self.para("– " + x[1], indent=1.25)

    # ---------- 3 Сетка, 4 Нагрузки ----------
    def sec3_4(self):
        me, t, c = self.p.raw["mesh"], self.p.tank, self.c
        self.h1("3 Сетка конструкции", new_page=True)
        beams = ", для элементов балки BEAM188" if me.get("beams") else ""
        self.body(f"Количество узлов сетки после ее разбиения: {me['nodes']} шт., количество конечных элементов: {me['elements']} шт. "
                  f"Качество сетки конструктивных элементов конструкции показано на рисунках 2 и 3. Размер элементов составляет {fmt(me['element_size_mm'], 0)} мм. "
                  f"Для элементов стенки используется элемент {me['element_type']}{beams}.")
        self.figure("fig02_mesh_wall_bottom", "Сетка элементов стенки и днища РВС", width_cm=16.5)
        self.figure("fig03_mesh_wall_roof", "Сетка элементов стенки и кровли РВС", width_cm=16.5)
        self.h1("4 Граничные условия и нагрузки", new_page=True)
        self.body("На конструкцию действует:")
        if self.rd:
            self.load_lines_ref()
        else:
            la = self.p.raw.get("loads_ansys")
            self.bullets(self.describe_loads(la) if la else self.legacy_load_bullets())
        self.body("Допущения, принятые в модели:")
        for a in T.ASSUMPTIONS:
            self.body(a)
        self.figure("fig04_loads", "Нагрузки", width_cm=16.5)
        self.figure("fig05_hydrostatic", "Гидростатическое давление", width_cm=16.5)

    def load_lines_ref(self):
        """Нагрузки раздела 4 формулами с числами (как в эталоне), из данных заказчика - те же числа ушли в расчёт Ansys."""
        from .loads.rvs_loads import CE
        t, rd = self.p.tank, self.rd
        G = 9.81
        L = lambda s: self.para("- " + s)
        rho = rd.get("product_density_kg_m3") or t.get("product_density")
        fill = (rd.get("fill_level_mm") or t.get("fill_level") or 0) / 1000
        L(f"гидростатическое давление продукта плотностью {fmt(rho, 0)} кг/м3, высота налива продукта {fmt(fill)} м.")
        L("вес металлоконструкций стенки, задан с помощью приложения ускорения свободного падения,")
        r = float(t["diameter"]) / 2000
        if rd.get("snow_kpa") is not None:
            S = rd["snow_gf"] * CE * rd["snow_kpa"] * 1000
            L(f"расчетная снеговая нагрузка {g3(rd['snow_kpa'])} кПа, приложенная к верхней грани стенки РВС, расчетное значение снеговой нагрузки "
              "согласно ГОСТ Р 58622-2019, ф. 87:")
            self.formula(["S = γf∙ce∙Sg = ", f"{g3(rd['snow_gf'])}∙{g3(CE)}∙{fmt(rd['snow_kpa'] * 1000, 0)} = {fmt(S, 0)} Па;"])
            self.formula(["Fсн = ", ("r∙S", "2"), " = ", (f"{trim(r, 3)}∙{fmt(S, 0)}", "2"), f" = {fmt(r * S / 2, 0)} Н/м."])
        if rd.get("roof_mass_kg") is not None:
            eq = rd.get("roof_equipment_kg") or []
            gf, psi = rd["roof_gf"], rd["psi"]
            F = gf * (rd["roof_mass_kg"] + psi * sum(eq)) * G
            mob = "(" + "+".join(g3(x) for x in eq) + ")" if len(eq) > 1 else (g3(eq[0]) if eq else "0")
            L("Нагрузка от веса крыши, оборудования и площадок обслуживания:")
            self.formula([f"F = γf∙(Mкр + ψ∙Mоб)∙g = {g3(gf)}∙({g3(rd['roof_mass_kg'])}+{g3(psi)}∙{mob})∙9,81 = {fmt(F, 1)} Н."])
        pk = rd.get("overpressure_kpa") or 0.0
        L("Расчетное избыточное давление Ри = γизб∙Ризб")
        self.formula([f"Ри = {g3(rd.get('overp_gf') or 1.2)}∙{fmt(pk * 1000, 0)} = {fmt((rd.get('overp_gf') or 1.2) * pk * 1000, 0)} Па"])
        ins = f" {fmt(rd['ins_thickness_mm'], 0)} мм" if rd.get("ins_thickness_mm") else ""
        if rd.get("wall_ins_kg"):
            L(f"Вес тепловой изоляции стенки{ins}, задается путем приложения силы ко всей площади стенки РВС, "
              f"Fизол ст = {g3(rd['wall_ins_kg'])}∙9,81 = {fmt(rd['wall_ins_kg'] * G, 0)} Н")
        if rd.get("roof_ins_kg"):
            L(f"Вес тепловой изоляции крыши{ins}, приложенный к верхней грани стенки РВС, Fизол кр = {g3(rd['roof_ins_kg'])}∙9,81 = {fmt(rd['roof_ins_kg'] * G, 0)} Н")
        if rd.get("wind_kpa") is not None:
            L(f"ветровая нагрузка (расчет устойчивости пустого резервуара): расчетное ветровое давление w0 = {g3(rd['wind_kpa'])} кПа, тип местности "
              f"{cyr(rd.get('terrain') or '')}, изменение по высоте – по СП 20.13330.2016;")
        L("Закрепление модели происходит по днищу путем задания ограничения по всем осям.")

    # ---------- нагрузки и свойства объекта ----------
    AXIS = {"X": "оси X (вертикаль)", "Y": "оси Y", "Z": "оси Z"}
    DIRS = {"NegativeXAxis": "по оси −X, вертикально вниз", "PositiveXAxis": "по оси +X",
            "NegativeYAxis": "по оси −Y", "PositiveYAxis": "по оси +Y",
            "NegativeZAxis": "по оси −Z", "PositiveZAxis": "по оси +Z"}

    def describe_loads(self, la: dict) -> list[str]:
        """Список нагрузок из проекта Ansys: имя, тип, значение, направление (если данных заказчика нет)."""
        labels, out = la.get("labels", {}), []
        for x in la["items"]:
            nm, t = x["name"], x["type"]
            lab = f" ({labels[nm]})" if nm in labels else ""
            if t == "EarthGravity":
                out.append(f"вес конструкций: ускорение свободного падения {fmt(x['g'], 5)} м/с2 ({self.DIRS.get(x['direction'], x['direction'])});")
            elif t == "HydrostaticPressure" and not x["suppressed"]:
                out.append(f"гидростатическое давление продукта плотностью {fmt(x['density'], 0)} кг/м3, уровень налива {fmt(x['level_m'])} м "
                           f"(ускорение {fmt(x['accel'], 1)} м/с2);")
            elif t == "HydrostaticPressure":
                out.append("гидростатическое давление в данном расчёте не учитывается;")
            elif t == "Force":
                out.append(f"«{nm}»{lab}: сила {fmt(abs(x['value']), 1)} Н, направление вдоль {self.AXIS[x['axis']]};")
            elif t == "LinePressure":
                out.append(f"«{nm}»{lab}: линейная нагрузка {fmt(abs(x['value']), 0)} Н/м по кромке стенки;")
            elif t == "Pressure":
                out.append(f"«{nm}»{lab}: давление {fmt(abs(x['value']), 0)} Па" + (" (в данном расчёте не учитывается);" if x["suppressed"] else ";"))
            elif t == "FixedSupport":
                out.append("закрепление модели по днищу: ограничение перемещений по всем осям;")
        return out

    def object_rows_ref(self) -> list[list]:
        """Таблица 1 как в эталоне ТО-019-26 (18 строк) по данным заказчика."""
        t, rd = self.p.tank, self.rd
        rho = rd.get("product_density_kg_m3") or t.get("product_density")
        fill = rd.get("fill_level_mm") or t.get("fill_level")
        eq = rd.get("roof_equipment_kg") or []
        rows = [["Высота стенки резервуара, мм", t["wall_height"]], ["Радиус стенки резервуара, мм", fmt(float(t["diameter"]) / 2, 0)],
                ["Радиус сферической крыши, мм", fmt(rd.get("roof_radius_mm") or t.get("roof_radius") or 0, 0)],
                ["Наименование хранимого продукта", rd.get("product") or t.get("product") or ""],
                ["Плотность продукта, кг/м3", fmt(rho, 0) if rho else ""], ["Расчетный уровень налива продукта, мм", fmt(fill, 0) if fill else ""],
                ["Нормативное внутреннее избыточное давление, кПа", g3(rd.get("overpressure_kpa") or 0)],
                ["Коэффициент надежности по нагрузке", g3(rd.get("load_factor"))],
                ["Собственный фактический вес крыши, оборудования, площадок и ограждений на крыше, кг", g3((rd.get("roof_mass_kg") or 0) + sum(eq))],
                ["Коэффициент надежности по нагрузке", g3(rd.get("roof_gf"))],
                ["Нагрузка от веса теплоизоляции крыши, кг", g3(rd.get("roof_ins_kg") or 0)],
                ["Нагрузка от веса теплоизоляции стенки, кг", g3(rd.get("wall_ins_kg") or 0)],
                ["Нормативный вес от снегового покрова, кПа", g3(rd.get("snow_kpa"))],
                ["Коэффициент надежности по снеговой нагрузке", g3(rd.get("snow_gf"))],
                ["Учет сдувания снега с крыши", rd.get("snow_blowoff") or ""], ["Тип местности", cyr(rd.get("terrain") or "")],
                ["Расчетное ветровое давление, кПа", g3(rd.get("wind_kpa"))],
                ["Класс опасности резервуара по ГОСТ 31385-2023", rd.get("hazard_class") or t.get("hazard_class") or ""]]
        return rows

    def object_rows(self) -> list[list]:
        t, L = self.p.tank, self.p.loads
        la = self.p.raw.get("loads_ansys")
        rows = [["Высота стенки резервуара, мм", t["wall_height"]], ["Радиус стенки резервуара, мм", fmt(t["diameter"] / 2, 0)]]
        if t.get("roof_radius"): rows.append(["Радиус сферической крыши, мм", t["roof_radius"]])
        if t.get("product"): rows.append(["Наименование хранимого продукта", t["product"]])
        if t.get("product_density"): rows.append(["Плотность продукта, кг/м3", fmt(t["product_density"], 0)])
        if t.get("fill_level"): rows.append(["Расчетный уровень налива продукта, мм", t["fill_level"]])
        if la:
            labels = la.get("labels", {})
            for x in la["items"]:
                if x["type"] == "Force":
                    rows.append([f"{x['name']}" + (f" – {labels[x['name']]}" if x['name'] in labels else "") + ", Н", fmt(abs(x["value"]), 1)])
                elif x["type"] == "LinePressure":
                    rows.append([f"{x['name']}" + (f" – {labels[x['name']]}" if x['name'] in labels else "") + ", Н/м", fmt(abs(x["value"]), 0)])
                elif x["type"] == "Pressure":
                    rows.append([f"{x['name']}" + (f" – {labels[x['name']]}" if x['name'] in labels else "") + ", Па", fmt(abs(x["value"]), 0)])
        elif L:
            rows += [["Нормативное внутреннее избыточное давление, кПа", fmt(L["overpressure_kpa"], 0)],
                     ["Коэффициент надежности по нагрузке", fmt(L["load_factor"])],
                     ["Собственный фактический вес крыши, оборудования и ограждений, кг", fmt(L["roof_mass_kg"] + sum(L["equipment_masses_kg"]), 1)],
                     ["Коэффициент надежности по нагрузке (крыша)", fmt(L["roof_load_factor"], 2)],
                     ["Нагрузка от веса теплоизоляции крыши, кг", L["roof_insulation_kg"]],
                     ["Нагрузка от веса теплоизоляции стенки, кг", L["wall_insulation_kg"]],
                     ["Нормативный вес от снегового покрова, кПа", fmt(L["snow_kpa"])],
                     ["Коэффициент надежности по снеговой нагрузке", fmt(L["snow_factor"])],
                     ["Учет сдувания снега с крыши", L["snow_blowoff"]], ["Тип местности", L["terrain"]],
                     ["Расчетное ветровое давление, кПа", fmt(L["wind_kpa"])]]
        for lab, val in (t.get("extra_rows") or []):
            rows.append([lab, val])
        if t.get("hazard_class"): rows.append(["Класс опасности резервуара по ГОСТ 31385", t["hazard_class"]])
        return rows

    def legacy_load_bullets(self) -> list[str]:
        t, L, c = self.p.tank, self.p.loads, self.c
        return [
            f"гидростатическое давление продукта плотностью {t['product_density']} кг/м3, высота налива {fmt(t['fill_level']/1000)} м;",
            "вес металлоконструкций стенки, задан приложением ускорения свободного падения;",
            f"расчетная снеговая нагрузка {fmt(L['snow_kpa'])} кПа, приложенная к верхней грани стенки: Fсн = {L['snow_linear_n_per_m']} Н/м;",
            f"нагрузка от веса крыши, оборудования и площадок: F = γf·(Mкр + ψ·Моб)·g = {fmt(L['roof_load_factor'],2)}·({L['roof_mass_kg']}+{fmt(L['combination_factor'],2)}·({' + '.join(str(x) for x in L['equipment_masses_kg'])}))·9,81 = {fmt(c['roof_force'],1)} Н;",
            f"расчетное избыточное давление Ри = {fmt(L['overpressure_factor'])}·{fmt(L['overpressure_kpa'],0)} = {fmt(L['overpressure_factor']*L['overpressure_kpa']*1000,0)} Па;",
            f"вес тепловой изоляции стенки, Fизол ст = {fmt(c['wall_ins_n'],0)} Н; вес тепловой изоляции крыши, Fизол кр = {fmt(c['roof_ins_n'],0)} Н;",
            "закрепление модели – по днищу, ограничение по всем осям."]

    # ---------- 5 Прочность, 6 Устойчивость ----------
    def strength_table(self, metric: str, key: str, limit: float, limit_label: str, ok_key: str, belts=None):
        """Таблица по поясам как в эталоне: двухуровневая шапка, строка номеров столбцов 1…5, без подписи «Таблица N»."""
        t = self.d.add_table(rows=3, cols=5)
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        h0, h1, h2 = t.rows
        self._merge_text(h0.cells[0], h1.cells[0], "№ п/п", size=11)
        self._merge_text(h0.cells[1], h1.cells[1], "Конструктивный элемент резервуара", size=11)
        self._merge_text(h0.cells[2], h0.cells[3], f"{metric}, МПа", size=11)
        self._merge_text(h0.cells[4], h1.cells[4], "Оценка выполнения критерия прочности", size=11)
        self._cell(h1.cells[2], "Максимальное расчетное значение напряжений", bold=True, align=CENTER, size=11)
        self._cell(h1.cells[3], f"Допустимое значение напряжений, {limit_label}", bold=True, align=CENTER, size=11)
        for i in range(5):
            self._cell(h2.cells[i], str(i + 1), align=CENTER, size=11)
        for r in (h0, h1, h2):
            self._repeat_header(r)
        for b in (belts or self.c["belts"]):
            cells = t.add_row().cells
            for i, v in enumerate([b["belt"], f"{b['belt']} пояс", fmt(b[key]), g3(round(limit, 1)), "Выполнен" if b[ok_key] else "Не выполнен"]):
                self._cell(cells[i], str(v), align=CENTER, red=(v == "Не выполнен"), size=11)
        t.autofit = False
        for r in t.rows:
            for i, w in enumerate([1.3, 4.0, 4.0, 4.0, 3.6]):
                r.cells[i].width = Cm(w)
        self.para(indent=0)
        return t

    def sec5_6(self):
        c, t, m = self.c, self.p.tank, self.p.material
        lv = fmt((t.get("fill_level") or 0) / 1000)
        self.h1("5 Результаты расчета на прочность", new_page=True)
        self.h1(f"5.1 Результаты расчета НДС РВС при наливе {lv} м с отклонениями стенки для заполненного резервуара")
        self.body(f"Обеспечение прочности поясов основной стенки резервуара {t['name']} в соответствии с ГОСТ 58622-2019 оценивается по выполнению следующих критериев:")
        self.para(f"– Срединные (мембранные) напряжения в кольцевом направлении σm в поясе основной стенки должны удовлетворять условию: σm ≤ [σ], "
                  f"где [σ] = σ02 / 1,5 – номинальное допускаемое напряжение для элементов оболочки резервуара, МПа; σ02 – условный предел текучести стали "
                  f"(σ02 = {g3(m['yield_mpa'])} МПа для стали {m['name']}" + (f" по {m['standard']}" if (m.get("standard") or "").strip() else "") + ").", indent=1.25)
        self.para("– Поверхностные (фибровые) напряжения в кольцевом направлении σф в поясе основной стенки должны удовлетворять условию: σф ≤ 3[σ].", indent=1.25)
        self.para("– Эквивалентные по Мизесу поверхностные напряжения σе в поясе основной стенки должны удовлетворять условию: σe ≤ 3[σ].", indent=1.25)
        self.body(f"[σ] = {g3(m['yield_mpa'])} / 1,5 = {fmt(c['allow'])} МПа; 3[σ] = {fmt(c['allow3'], 0)} МПа.")
        self.figure("fig06_eq_stress", "Эквивалентные поверхностные напряжения в стенке РВС", width_cm=16.5)
        self.figure("fig07_membrane", "Срединные (мембранные) напряжения в кольцевом направлении", width_cm=16.5)
        self.figure("fig08_fiber", "Поверхностные (фибровые) напряжения в кольцевом направлении", width_cm=16.5)
        self.figure("fig09_deform", "Деформации стенки РВС", width_cm=16.5)
        self.strength_table("Первые главные поверхностные (фибровые) напряжения σф", "fiber", c["allow3"], "3[σ]", "ok_fiber")
        self.strength_table("Эквивалентные по Мизесу поверхностные напряжения σе", "equivalent", c["allow3"], "3[σ]", "ok_eq")
        self.strength_table("Первые главные срединные (мембранные) напряжения σм", "membrane", c["allow"], "[σ]", "ok_mem")
        self.h1(f"5.2 Результаты расчета НДС РВС при наливе {lv} м с отклонениями стенки для пустого резервуара", new_page=True)
        self.figure("fig10_empty_eq", "Эквивалентные напряжения в стенке РВС", width_cm=16.5)
        self.figure("fig11_empty_mem", "Мембранные (срединные) напряжения в стенке РВС", width_cm=16.5)
        self.figure("fig12_empty_fiber", "Фибровые (поверхностные) напряжения в стенке РВС, максимальное и минимальное главное напряжение", width_cm=16.5)
        self.figure("fig13_empty_deform", "Деформации стенки РВС", width_cm=16.5)
        self.h1("6 Результаты расчета РВС на устойчивость", new_page=True)
        self.h1("6.1 Результаты расчета РВС на устойчивость искривленного резервуара без гидростатического давления и с ветровой нагрузкой")
        for i, s in enumerate(self.c["stability"], 1):
            k = g3(round(s["k"], 1)) if s["k"] > 0 else "нет данных"
            self.figure(f"fig_stab{i}", f"Коэффициент запаса устойчивости Fкр/F = k = {k}, {s['case'][0].lower() + s['case'][1:]}", width_cm=16.0)

    def _allowed_fill(self) -> str | None:
        res, c = self.p.raw["results"], self.c
        if res.get("fill_series") and len(res["fill_series"]) >= 3:
            from .fillcalc import allowable_fill
            fa = allowable_fill(res["fill_series"], c["allow"])
            return ("не менее " if fa["unlimited"] else "") + fmt(fa["level_m"])
        return fmt(res["max_fill_m"]) if c["strength_ok"] and res.get("max_fill_m") else None

    def sec8(self):
        c, t, m, res = self.c, self.p.tank, self.p.material, self.p.raw["results"]
        n = len(t["belts_thickness"])
        self.h1("7 Выводы и рекомендации", new_page=True)
        rho = t.get("product_density") or "–"
        items = [f"Расчеты прочности стенки {t['name']} при наливе продукта с плотностью {fmt(float(rho), 0) if rho != '–' else rho} кг/м3 до расчетной высоты налива "
                 f"{fmt((t.get('fill_level') or 0) / 1000)} м, с учетом фактической формы стенки показали, что прочность стенки резервуара с учетом имеющихся "
                 "отклонений при действии расчетных нагрузок и отклонениях от вертикали образующих стенки на высоте каждого пояса заполненного резервуара "
                 + ("обеспечена." if c["strength_ok"] else "НЕ обеспечена.")]
        i = res.get("ideal")
        if i:
            ok_i = i.get("membrane", 0) <= c["allow"] and i.get("fiber", 0) <= c["allow3"] and i.get("equivalent", 0) <= c["allow3"]
            items.append(f"Максимальные эквивалентные напряжения в стенке идеального РВС составили {fmt(i['equivalent'])} МПа, мембранные напряжения {fmt(i['membrane'])} МПа, "
                         f"фибровые напряжения {fmt(i['fiber'])} МПа" + (f", максимальные перемещения {fmt(i['deform_mm'])} мм" if i.get("deform_mm") else "")
                         + f". Допустимые расчетные сопротивления согласно ГОСТ 31385-2023 для 1го пояса составили {int(c['R1'])} МПа, 2-{n} го пояса составили "
                         f"{fmt(c['R2'])} МПа. Прочность идеализированного РВС без отклонений " + ("обеспечена" if ok_i else "НЕ обеспечена")
                         + (f", коэффициент запаса устойчивости k = {fmt(i['k'])}." if i.get("k") else "."))
        dmax = res.get("max_deform_mm")
        items.append(f"Максимальные эквивалентные напряжения в стенке РВС с отклонениями от идеальной формы составили {fmt(c['max_eq'])} МПа, мембранные напряжения "
                     f"{fmt(c['max_mem'])} МПа, фибровые напряжения {fmt(c['max_fiber'])} МПа" + (f", максимальные перемещения {fmt(dmax)} мм" if dmax else "")
                     + ". Прочность РВС " + ("обеспечена." if c["strength_ok"] else "НЕ обеспечена."))
        if c["kmin"] is not None:
            smin = min((s for s in c["stability"] if s["k"] > 0), key=lambda s: s["k"])
            req = res["required_k"]
            items.append(f"Запас устойчивости стенки составил k = {g3(round(c['kmin'], 1))} для резервуара с отклонениями пустого РВС, "
                         f"{smin['case'][0].lower() + smin['case'][1:]}, что является " + ("достаточным" if c["stability_ok"] else "НЕДОСТАТОЧНЫМ")
                         + f" запасом устойчивости (требуется k ≥ {g3(req)}). Данный результат не зависит от высоты налива, поскольку считается пустой РВС под ветровой нагрузкой.")
        else:
            items.append("Коэффициент запаса устойчивости не получен (нет результатов расчета на устойчивость).")
        af = self._allowed_fill()
        items.append(f"Допустимый уровень налива продукта по условию прочности составил {af} м." if af else
                     "Для продолжения эксплуатации требуется снижение уровня налива или вывод резервуара в ремонт; допустимый уровень налива следует определить повторным расчетом.")
        for k, x in enumerate(items, 1):
            par = self.body(f"{k}. {x}")
            par.paragraph_format.keep_together = True

    def appendices(self):
        self.h1("Приложение 1 – Перечень использованной при техническом диагностировании нормативной технической и методической документации", new_page=True)
        for k, x in enumerate(T.NORMS_FULL, 1): self.body(f"{k}. {x}")
        self.h1("Приложение 2 – Перечень используемых терминов и определений", new_page=True)
        for term, dfn in T.TERMS:
            self.body(term); self.body(dfn)
        self.h1("Приложение 3 – Исходные данные на выполнение прочностного расчет РВС", new_page=True)
        self.body("1. Исходные данные")
        n = self.p.raw["survey"]["points_per_belt"]
        nb = len(self.p.tank["belts_thickness"])
        for title, dev in (("Таблица 1.1 – Отклонения стенки от вертикали полного резервуара.", self.p.survey_full),
                           ("Таблица 1.2 – Отклонения стенки от вертикали резервуара после опорожнения.", self.p.survey_empty)):
            self.para(title, align=WD_ALIGN_PARAGRAPH.LEFT, indent=0, keep_next=True, space_before=6, space_after=3)
            hdr = ["Образующая"] + [f"{ROMAN[b]}" for b in range(nb)]
            rows = [[i + 1] + [int(dev[b][i]) if float(dev[b][i]).is_integer() else fmt(dev[b][i]) for b in range(nb)] for i in range(n)]
            self.table(hdr, rows, size=9)
        self.body("2. Требование к моделированию геометрической модели:")
        for line in T.REQUIREMENTS: self.body(line)

    def build(self, out: Path) -> Path:
        self.header_footer(); self.title_page(); self.toc()
        self.sec1(); self.sec2(); self.sec3_4(); self.sec5_6(); self.sec8(); self.appendices()
        out.parent.mkdir(parents=True, exist_ok=True)
        self.d.save(out)
        return out


def g3(x) -> str:
    """Число как в эталоне: до 3 значащих после запятой без лишних нулей (7,45; 14,9; 0,95; 64156; 1571,8)."""
    if x is None or x == "":
        return ""
    x = float(x)
    if x == int(x):
        return str(int(x))
    if abs(x) >= 100:
        return f"{x:.1f}".rstrip("0").rstrip(".").replace(".", ",")
    return f"{x:.3g}".replace(".", ",")


def trim(x: float, nd: int = 3) -> str:
    """14.25 -> «14,25» (не более nd знаков после запятой, без лишних нулей)."""
    return f"{float(x):.{nd}f}".rstrip("0").rstrip(".").replace(".", ",")


def cyr(s: str) -> str:
    """Тип местности латиницей (A, B, C) → кириллицей, как в эталоне."""
    return str(s).translate(str.maketrans("ABC", "АВС"))


def build_report(p: Project, out: Path) -> tuple[Path, list[str]]:
    rep = Report(p)
    path = rep.build(out)
    return path, rep.missing
