import json
"""Тесты на обезличенном примере examples/demo (данные получены из реального проекта Ansys)."""
from pathlib import Path
import openpyxl
from ansys_report.project import load_project
from ansys_report.calc import calc_all
from ansys_report.excel_gen import build_excel
from ansys_report.word_gen import build_report

YAML = Path(__file__).parent.parent / "examples/demo/project.yaml"


def test_design_numbers():
    c = calc_all(load_project(YAML))
    assert round(c["R1"], 2) == 206.35 and round(c["R2"], 2) == 235.83 and round(c["Rutor"], 2) == 353.74
    assert round(c["E2"], 1) == 743.7 and round(c["allow"], 1) == 216.7
    assert c["strength_ok"] and c["kmin"] > 10 and c["stability_ok"]


def test_excel_and_word_build(tmp_path):
    p = load_project(YAML)
    x = build_excel(p, tmp_path / "a.xlsx")
    wb = openpyxl.load_workbook(x)
    assert wb.sheetnames == ["Т-001 полный", "Т-001 пустой", "Проверка"]
    docx_path, missing = build_report(p, tmp_path / "a.docx")
    assert docx_path.stat().st_size > 10_000 and "fig01_general" in missing



def test_gui_core(tmp_path):
    from ansys_report.gui import build_all
    msgs = []
    files = build_all(YAML, tmp_path, msgs.append)
    assert len(files) == 5 and all(f.exists() for f in files) and any("Word" in m for m in msgs)


def test_apdl_shift_matches_survey():
    """В узлах замера (пояс k, точка j) сдвиг должен равняться (полный − пустой)."""
    import math
    from ansys_report.apdl_gen import delta_table, radial_shift_mm, build_apdl
    p = load_project(Path(__file__).parent.parent / "examples/demo/project.yaml")
    tbl = delta_table(p)
    s, h = p.raw["survey"], p.tank["belt_height"] / 1000
    n = s["points_per_belt"]
    for k in (1, 5, 12):
        for j in (0, 3, 7, 14):
            a = math.radians(s["start_angle"] - 360 / n * j)
            R = 14.25
            got = radial_shift_mm(p, tbl, k * h, R * math.cos(a), R * math.sin(a))
            want = p.survey_full[k - 1][j] - p.survey_empty[k - 1][j]
            assert abs(got - want) < 1e-6, (k, j, got, want)
    assert abs(radial_shift_mm(p, tbl, 0.0, 14.25, 0.0)) < 1e-9     # днище не двигается
    txt = build_apdl(p)
    assert "NMODIF" in txt and "*DIM,DLT,ARRAY,13,16" in txt and txt.isascii()


def test_report_from_ansys_model_data(tmp_path):
    """Режим «всё из Ansys»: геометрия, сетка, нагрузки и нивелировка «пустой» берутся из model_data.json."""
    import docx
    y = Path(__file__).parent.parent / "examples/demo/project.yaml"
    p = load_project(y)
    assert p.tank["belts_thickness"][:3] == [16.0, 14.0, 13.0] and p.tank["wall_height"] == 17880
    assert p.raw["mesh"]["nodes"] == 60694 and p.tank["fill_level"] == 17200
    assert len(p.survey_empty) == 12 and len(p.survey_empty[0]) == 15
    path, _ = build_report(p, tmp_path / "a.docx")
    text = "\n".join(x.text for x in docx.Document(path).paragraphs)
    assert "60694" in text and "734643,9" in text and "62907" not in text


def test_wizard_project_folder(tmp_path):
    """Папка выгрузки + анкета из мастера → project.yaml → Excel/Word (без запуска Ansys)."""
    import shutil
    from ansys_report.ansys_connect import make_project_folder
    from ansys_report.gui import build_all
    src = Path(__file__).resolve().parent.parent / "examples" / "demo"
    data = tmp_path / "data"
    shutil.copytree(src, data, ignore=shutil.ignore_patterns("project.yaml", "survey_full.csv", "images"))
    yaml_path = make_project_folder(data, {"number": "ТО-777-26", "year": 2026, "tag": "Т-777", "tank_name": "РВС-1",
                                           "roof_radius": 34200, "executor_name": "Иванов И.И."})
    files = build_all(yaml_path, tmp_path / "out", log=lambda m: None)
    assert any(f.suffix == ".docx" for f in files) and any(f.suffix == ".xlsx" for f in files)


def test_roles_and_custom_belt_regex(tmp_path):
    """Соответствия: роли расчётов выбираются по настройке/имени, имена поясов — по регулярному выражению."""
    import json
    import re
    import shutil
    from ansys_report.settings import ModelSettings, pick_roles
    from ansys_report.ansys_connect import make_project_folder
    from ansys_report import ansys_import as ai
    names = ["Static Structural", "Static Structural 2", "Eigenvalue Buckling", "Static Structural Full"]
    assert pick_roles(names, ModelSettings()) == ("Static Structural Full", "Static Structural 2")
    assert pick_roles(names, ModelSettings(full_analysis="Static Structural", empty_analysis="Static Structural 2")) == \
        ("Static Structural", "Static Structural 2")
    assert pick_roles(["Eigenvalue Buckling"], ModelSettings()) == (None, None)
    src = Path(__file__).resolve().parent.parent / "examples" / "demo"
    data = tmp_path / "d"
    shutil.copytree(src, data, ignore=shutil.ignore_patterns("project.yaml", "survey_full.csv", "images"))
    # пояса названы иначе: «Belt_01…Belt_12»
    md = json.loads((data / "model_data.json").read_text(encoding="utf-8"))
    for b in md["bodies"]:
        m = re.match(r"^p\s*(\d+)$", b["name"])
        if m:
            b["name"] = "Belt_%02d" % int(m.group(1))
    md["belt_regex"] = r"^Belt_(\d+)$"
    (data / "model_data.json").write_text(json.dumps(md), encoding="utf-8")
    assert len(ai.belt_bodies(md)) == 12
    st = ModelSettings(buckling_labels={"Eigenvalue Buckling": "Ветер A", "Eigenvalue Buckling 2": "Ветер B"})
    p = load_project(make_project_folder(data, {"number": "N", "year": 2026, "tag": "T-1", "roof_radius": 1}, st))
    assert p.has_results and p.results_empty and "Ветер A" in p.stability[0]["case"]


def _renamed_demo(tmp_path):
    """Копия примера с другими именами: пояса «Ring_01…», расчёты «Strength …/Buckle …» (проверка автоматики соответствий)."""
    import json
    import re
    import shutil
    src = Path(__file__).resolve().parent.parent / "examples" / "demo"
    data = tmp_path / "renamed"
    shutil.copytree(src, data, ignore=shutil.ignore_patterns("project.yaml", "survey_full.csv", "images"))
    ren = {"Static Structural Full": "Strength Full", "Static Structural 2": "Strength Empty",
           "Static Structural": "Strength Wind", "Eigenvalue Buckling 2": "Buckle Z", "Eigenvalue Buckling": "Buckle Y"}

    def new(name):
        for old in sorted(ren, key=len, reverse=True):      # сначала длинные имена
            if name == old:
                return ren[old]
        return name

    md = json.loads((data / "model_data.json").read_text(encoding="utf-8"))
    for b in md["bodies"]:
        m = re.match(r"^p\s*(\d+)$", b["name"])
        if m:
            b["name"] = "Ring_%02d" % int(m.group(1))
    md["belt_regex"] = r"^Ring_(\d+)$"
    for a in md["analyses"]:
        a["name"] = new(a["name"])
    (data / "model_data.json").write_text(json.dumps(md), encoding="utf-8")
    for f in list(data.glob("*_belts.csv")):
        base = f.stem[: -len("_belts")].replace("_", " ")
        f.rename(data / (new(base).replace(" ", "_") + "_belts.csv"))
    rows = (data / "stability.csv").read_text(encoding="utf-8").splitlines()
    (data / "stability.csv").write_text("\n".join([rows[0]] + [new(r.split(",")[0]) + "," + ",".join(r.split(",")[1:]) for r in rows[1:]]),
                                         encoding="utf-8")
    img = data / "images"
    img.mkdir()
    from docx.shared import Pt  # noqa: F401  (python-docx уже нужен)
    import struct, zlib
    png = (b"\x89PNG\r\n\x1a\n" + struct.pack(">I", 13) + b"IHDR" + struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0) +
           struct.pack(">I", zlib.crc32(b"IHDR" + struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))) + struct.pack(">I", 12) + b"IDAT" +
           zlib.compress(b"\x00\x00\x00\x00") + struct.pack(">I", zlib.crc32(b"IDAT" + zlib.compress(b"\x00\x00\x00\x00"))) +
           struct.pack(">I", 0) + b"IEND" + struct.pack(">I", zlib.crc32(b"IEND")))
    for n in ("Strength_Full_Equivalent_Stress", "Strength_Empty_Equivalent_Stress", "Buckle_Y_Total_Deformation",
              "Buckle_Z_Total_Deformation_2"):
        (img / (n + ".png")).write_bytes(png)
    return data


def test_renamed_project_gives_same_report(tmp_path):
    from ansys_report.settings import ModelSettings
    from ansys_report.ansys_connect import make_project_folder
    from ansys_report.gui import build_all
    orig = calc_all(load_project(YAML))
    data = _renamed_demo(tmp_path)
    st = ModelSettings(belt_regex=r"^Ring_(\d+)$", static_prefix="Strength", buckling_prefix="Buckle",
                       buckling_labels={"Buckle Y": "Направление ветра Y", "Buckle Z": "Направление ветра Z"})
    y = make_project_folder(data, {"number": "N", "year": 2026, "tag": "T-2", "roof_radius": 1}, st)
    p = load_project(y)
    c = calc_all(p)
    assert c["max_eq"] == orig["max_eq"] and c["max_fiber"] == orig["max_fiber"] and c["kmin"] == orig["kmin"]
    assert [s["case"] for s in c["stability"]][0].startswith("Направление ветра Y")
    assert p.raw["ansys"]["analysis_full"] == "Strength Full"
    assert p.raw["images"]["fig06_eq_stress"] == "Strength_Full_Equivalent_Stress.png"
    assert p.raw["images"]["fig10_empty_eq"] == "Strength_Empty_Equivalent_Stress.png"
    assert p.raw["images"]["fig_stab1"] == "Buckle_Y_Total_Deformation.png"
    files = build_all(y, tmp_path / "out", log=lambda m: None)
    assert any(f.suffix == ".docx" for f in files)


def test_web_ui_api(tmp_path, monkeypatch):
    """Веб-интерфейс: состояние, сборка из готовой выгрузки (без Ansys), статус, ответ страницы."""
    import json
    import shutil
    import threading
    import time
    import urllib.request
    from http.server import ThreadingHTTPServer
    from ansys_report import webui
    monkeypatch.setattr("ansys_report.wizard.SETTINGS", tmp_path / "settings.json")
    job, state = webui.Job(), {"last_ping": time.time()}
    srv = ThreadingHTTPServer(("127.0.0.1", 0), webui.make_handler(job, state))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_address[1]}"

    def call(path, body=None):
        req = urllib.request.Request(base + path, data=None if body is None else json.dumps(body).encode(),
                                     headers={"Content-Type": "application/json"})
        return json.loads(urllib.request.urlopen(req).read())

    assert b"Ansys Report" in urllib.request.urlopen(base + "/").read()
    assert call("/api/state")["model"]["belt_regex"].startswith("^p")
    out = tmp_path / "out"
    data = out / "Report_T-9" / "_ansys_data"
    shutil.copytree(Path(__file__).resolve().parent.parent / "examples" / "demo", data,
                    ignore=shutil.ignore_patterns("project.yaml", "survey_full.csv", "images"))
    anketa = {"number": "N-1", "year": "2026", "tag": "T-9", "roof_radius": "34200", "executor_name": "X", "site": "S"}
    assert call("/api/build", {"project": "C:/x/Tank.wbpj", "out": str(out), "anketa": anketa, "solve": False,
                               "reuse": True, "model": {}})["ok"]
    for _ in range(100):
        st = call("/api/status?since=0")
        if st["done"]:
            break
        time.sleep(0.2)
    assert st["done"] and not st["error"] and any(f.endswith(".docx") for f in st["files"])
    srv.shutdown()


_ENGDATA = """<?xml version="1.0" encoding="utf-8"?>
<EngineeringData version="2026 R1"><Materials><MatML_Doc>
<Material><BulkDetails><Name>Structural Steel</Name>
<PropertyData property="pr0"><Data format="string">-</Data><ParameterValue parameter="pa0" format="float"><Data format="string">7850</Data></ParameterValue></PropertyData>
<PropertyData property="pr1"><Data format="string">-</Data><ParameterValue parameter="pa1" format="float"><Data format="string">2E+11</Data></ParameterValue>
<ParameterValue parameter="pa2" format="float"><Data format="string">0.3</Data></ParameterValue></PropertyData>
<PropertyData property="pr2"><Data format="string">-</Data><ParameterValue parameter="pa3" format="float"><Data format="string">250000000</Data></ParameterValue></PropertyData>
<PropertyData property="pr3"><Data format="string">-</Data><ParameterValue parameter="pa4" format="float"><Data format="string">460000000</Data></ParameterValue></PropertyData>
</BulkDetails></Material>
<Metadata><ParameterDetails id="pa0"><Name>Density</Name></ParameterDetails>
<ParameterDetails id="pa1"><Name>Young's Modulus</Name></ParameterDetails>
<ParameterDetails id="pa2"><Name>Poisson's Ratio</Name></ParameterDetails>
<ParameterDetails id="pa3"><Name>Tensile Yield Strength</Name></ParameterDetails>
<ParameterDetails id="pa4"><Name>Tensile Ultimate Strength</Name></ParameterDetails>
<PropertyDetails id="pr0"><Name>Density</Name></PropertyDetails></Metadata>
</MatML_Doc></Materials></EngineeringData>
"""


def test_material_from_ansys_engineering_data(tmp_path):
    import shutil
    from ansys_report import ansys_import as ai
    f = tmp_path / "EngineeringData.xml"
    f.write_text(_ENGDATA, encoding="utf-8")
    mats = ai.read_engineering_data(f)
    assert ai.material_summary(mats["Structural Steel"]) == {"density": 7850.0, "young_modulus_mpa": 200000.0,
                                                           "poisson": 0.3, "yield_mpa": 250.0, "ultimate_mpa": 460.0}
    src = Path(__file__).resolve().parent.parent / "examples" / "demo"
    d = tmp_path / "p"
    shutil.copytree(src, d, ignore=shutil.ignore_patterns("project.yaml", "survey_full.csv", "images"))
    shutil.copy(f, d / "engineering_data.xml")
    from ansys_report.ansys_connect import make_project_folder
    p = load_project(make_project_folder(d, {"number": "N", "year": 2026, "tag": "T-3", "roof_radius": 1}))
    assert p.material["yield_mpa"] == 250 and p.material["young_modulus_mpa"] == 200000
    assert isinstance(p.material["yield_mpa"], int) and isinstance(p.material["density"], int)
    assert p.material["name"] == "Structural Steel" and p.material["source"] == "Ansys"
    p2 = load_project(make_project_folder(d, {"number": "N", "year": 2026, "tag": "T-3", "roof_radius": 1, "steel_name": "09Г2С"}))
    assert p2.material["name"] == "09Г2С"
    c = calc_all(p)
    assert abs(c["allow"] - 250 / 1.5) < 0.01
    assert abs(c["allow15"] - 250) < 0.01
    assert c["strength_ok"] == all(b["ok_mem"] and b["ok_fiber"] and b["ok_eq"] for b in c["belts"])
    assert all(b["ok_mem"] == (b["membrane"] <= c["allow"]) for b in c["belts"])         # мембранные: [σ], как в эталоне ТО-019-26
    assert len(c["dev_check"]) == len(c["belts"]) and c["dev_check"][0]["lim2"] == c["dev_check"][0]["lim1"] + 30


def test_full_shape_apdl_from_survey_csv(tmp_path):
    """Замеры «полного» (CSV) + выгрузка модели -> макрос APDL с ненулевыми сдвигами; шаблон замеров читается обратно."""
    import shutil
    from ansys_report.apdl_gen import build_apdl
    from ansys_report.ansys_connect import make_project_folder
    from ansys_report.project import write_survey_csv
    src = Path(__file__).resolve().parent.parent / "examples" / "demo"
    d = tmp_path / "d"
    shutil.copytree(src, d, ignore=shutil.ignore_patterns("project.yaml", "images"))
    p = load_project(make_project_folder(d, {"number": "N", "year": 2026, "tag": "T-5", "roof_radius": 1}))
    txt = build_apdl(p)
    import re
    assert "NMODIF" in txt and re.search(r"^DLT\(\d+,\d+\) = ", txt, re.M) and txt.isascii()
    t = write_survey_csv(tmp_path / "tpl.csv", p.survey_empty)       # шаблон = «пустой»: после чтения сдвиги нулевые
    d2 = tmp_path / "d2"
    shutil.copytree(src, d2, ignore=shutil.ignore_patterns("project.yaml", "images"))
    shutil.copy(t, d2 / "survey_full.csv")
    p2 = load_project(make_project_folder(d2, {"number": "N", "year": 2026, "tag": "T-5", "roof_radius": 1}))
    assert not re.search(r"^DLT\(\d+,\d+\) = ", build_apdl(p2), re.M)


def test_universal_report_for_acoustic_project(tmp_path):
    """Проект без поясов (Modal Acoustics): универсальный отчёт собирается без ошибок."""
    import docx
    from ansys_report.universal import build_universal
    md = {"vertical_axis": "X", "belt_rings": [], "belt_regex": r"^p\s*(\d+)$", "node_count": 10, "element_count": 5,
          "bbox": {"xmax": 2.5, "xmin": -2.5, "rmax": 6.3},
          "mesh_object": {"ElementSize": "0.25 [m]", "ElementOrder": "Linear"},
          "bodies": [{"name": "Solid", "material": "Structural Steel", "thickness": "0 [m]", "dimension": "?", "suppressed": "False"}],
          "materials": [{"name": "Structural Steel", "props": {"AssignedBodies": "1"}}],
          "analyses": [{"name": "Modal Acoustics", "type": "Modal", "results": [], "children": [
              {"name": "Analysis Settings", "type": "ANSYSAnalysisSettings", "props": {"MaximumModesToFind": "6"}},
              {"name": "Physics Region", "type": "PhysicsRegion", "props": {"Acoustics": "True"}}]}]}
    (tmp_path / "model_data.json").write_text(json.dumps(md), encoding="utf-8")
    (tmp_path / "modes.csv").write_text("analysis,type,mode,value,unit\nModal Acoustics,Modal,1,12.5,Hz\n", encoding="utf-8")
    out, _ = build_universal(tmp_path, tmp_path / "r.docx", {"tank_name": "Башня"})
    text = "\n".join(p.text for p in docx.Document(out).paragraphs)
    assert "Modal Acoustics" in text and "5 Результаты расчёта" in text


def _points_project():
    """Синтетический проект замеров: 12 образующих × 13 отметок, наклон образующих, одна вмятина, три пояса толщины."""
    import math
    angles = [30.0 * i for i in range(12)]
    heights = [1.5 * k for k in range(13)]
    values = [[f"{(1 if i % 2 == 0 else -1) * z * (5 + i) * 0.6:.1f}".replace(".", ",") for i in range(12)] for z in heights]
    dent_x = [-0.6, -0.3, 0.0, 0.3, 0.6]
    dent_y = [-0.6, -0.3, 0.0, 0.3, 0.6]
    dv = [[("0" if (j in (0, 4) or i in (0, 4)) else ("-12" if (i == 2 and j == 2) else "-3")) for i in range(5)] for j in range(5)]
    return {"version": "1.0", "diameter": "28.5", "height": "18", "angles": angles, "heights": heights, "values": values,
            "bands": [["1.5", "13"]] + [["1.5", "12"]] * 2 + [["1.5", "11"]] * 9,
            "defects": [{"name": "Вмятина 1", "mode": "Дуга, м", "anchor": "20", "height": "6", "width": "1.2", "span_height": "1.2",
                         "nx": "5", "ny": "5", "x": dent_x, "y": dent_y, "values": dv}],
            "contour_threshold": "0.1", "wall_includes_defects": False}


def test_points_geometry_check_and_report(tmp_path):
    """Замеры → оценка по допускам H/200 и H/200+30 → CAD-оболочка → Word."""
    import docx
    from ansys_report.geometry import check
    from ansys_report.geometry.pipeline import build_from_points
    pr = _points_project()
    res = check.analyze(pr)
    assert res["n_points"] == 12 * 12 and res["over_a"] >= res["over_b"] >= 0
    b1 = res["belts"][0]
    assert abs(b1["lim_a"] - 7.5) < 1e-9 and abs(b1["lim_b"] - 37.5) < 1e-9          # H/200 и H/200 + 30 на первом поясе
    assert res["dents"][0]["depth_mm"] > 11
    f = tmp_path / "p.json"
    f.write_text(json.dumps(pr, ensure_ascii=False), encoding="utf-8")
    files = build_from_points(f, tmp_path / "out", {"tag": "T-77", "number": "N-1", "tank_name": "РВС-1"}, log=lambda s: None)
    names = [x.name for x in files]
    assert any(n.endswith(".docx") for n in names) and "Поверхность РВС.step" in names and "Поверхность РВС.igs" in names
    step = next(x for x in files if x.suffix == ".step").read_text(encoding="ascii")
    assert "B_SPLINE_SURFACE_WITH_KNOTS" in step and step.count("ADVANCED_FACE") == 3      # три участка толщины 13 / 12 / 11 мм
    text = "\n".join(p.text for p in docx.Document(next(x for x in files if x.suffix == ".docx")).paragraphs)
    assert "Оценка отклонений по точкам" in text and "5 Выводы" in text


def test_ideal_cache_key_follows_inputs(tmp_path):
    """Кэш идеального РВС: те же файлы пакета - тот же ключ, любое изменение - другой (считаем заново)."""
    from ansys_report.agent import _ideal_cache
    (tmp_path / "model.step").write_text("STEP-1", encoding="utf-8")
    (tmp_path / "loads_data.py").write_text("LOADS = {}", encoding="utf-8")
    k1 = _ideal_cache(tmp_path)
    assert _ideal_cache(tmp_path) == k1 and k1.suffix == ".json"
    (tmp_path / "model.step").write_text("ISO-10303-21;\nHEADER;\nFILE_NAME('m','2026-10-27T01:31:41');\nENDSEC;\nDATA;\n#1=POINT;", encoding="utf-8")
    k2 = _ideal_cache(tmp_path)
    (tmp_path / "model.step").write_text("ISO-10303-21;\nHEADER;\nFILE_NAME('m','2026-10-28T09:00:00');\nENDSEC;\nDATA;\n#1=POINT;", encoding="utf-8")
    assert _ideal_cache(tmp_path) == k2                     # другое время в заголовке STEP - тот же ключ
    (tmp_path / "loads_data.py").write_text("LOADS = {'snow': 1}", encoding="utf-8")
    assert _ideal_cache(tmp_path) != k2


def test_agent_lock_blocks_second_run(tmp_path, monkeypatch):
    """Второй помощник не трогает рабочую папку, пока жив первый; брошенная блокировка (процесс умер) снимается."""
    import os
    from ansys_report import agent
    monkeypatch.setattr(agent, "WORK", tmp_path)
    lock = agent._take_lock()
    assert lock is not None and lock.read_text() == str(os.getpid())
    assert agent._take_lock() is None                       # занята живым процессом (этим)
    lock.write_text("999999")                               # такого процесса нет
    assert agent._take_lock() is not None


def test_timing_table(tmp_path):
    from ansys_report.timing import read, table_md
    f = tmp_path / "timing.log"
    f.write_text("100.0\t\tАгент: старт\n110.0\tмодель\tСборка: сетка\n140.0\tмодель\tСборка: нагрузки\n150.0\t\tконец\n", encoding="utf-8")
    t = table_md(read(f))
    assert "Всего: 50 с" in t and "| модель | Сборка: сетка | 30.0 | 60.0 |" in t and "| модель | Сборка | 40.0 | 80.0 |" in t


def test_unsolved_ansys_results_stop_report(tmp_path):
    """Все пояса нулевые (сетка не построилась, расчёт не шёл) - ошибка с причиной из лога, а не отчёт с «Выполнен»."""
    import pytest
    from ansys_report.ansys_connect import ConnectError, _report_problems
    zero = "belt,fiber,equivalent,membrane\n1,0.00,0.00,0.00\n2,0.00,0.00,0.00\n"
    for n in ("AR__Эксплуатация_belts.csv", "AR__Налив_7.2_м_belts.csv"):
        (tmp_path / n).write_text(zero, encoding="utf-8")
    (tmp_path / "log.txt").write_text("  сообщение [Error]: One or more faces failed to mesh in the meshing engine.\n", encoding="utf-8")
    with pytest.raises(ConnectError, match="failed to mesh"):
        _report_problems(tmp_path, log=lambda s: None)
    (tmp_path / "AR__Эксплуатация_belts.csv").write_text("belt,fiber,equivalent,membrane\n1,5.0,6.0,3.0\n", encoding="utf-8")
    _report_problems(tmp_path, log=lambda s: None)            # хотя бы один случай решён - только предупреждение


def test_dent_zones_for_mesh_refinement():
    """Зона сгущения сетки у вмятины: дуга 20 м на R = 14,25 м -> угол 20/14,25 рад; радиус - полудиагональ участка
    с переходной зоной (±0,6 м + 2 шага по 0,3 м = ±1,2 м) плюс 200 мм."""
    import math
    import pytest
    from ansys_report.geometry.check import dent_zones
    z = dent_zones(_points_project())
    assert len(z) == 1
    th = 20.0 / 14.25
    assert z[0]["x_mm"] == pytest.approx(14250 * math.cos(th), abs=0.1) and z[0]["y_mm"] == pytest.approx(14250 * math.sin(th), abs=0.1)
    assert z[0]["z_mm"] == pytest.approx(6000.0) and z[0]["radius_mm"] == pytest.approx(1200 * math.sqrt(2) + 200, abs=0.1)


def test_report_with_ideal_tank(tmp_path):
    """Результаты идеального РВС попадают в выводы и в таблицу устойчивости."""
    import docx
    p = load_project(YAML)
    p.raw["results"]["ideal"] = {"equivalent": 111.1, "membrane": 99.9, "fiber": 88.8, "k": 2.5}
    path, _ = build_report(p, tmp_path / "i.docx")
    text = "\n".join(x.text for x in docx.Document(path).paragraphs)
    assert "идеального РВС" in text and "111,1" in text and "k = 2,5" in text


def test_report_with_fill_series(tmp_path):
    """Допустимая высота налива по серии уровней (критерий σм ≤ [σ]) - в выводах; отдельной таблицы, как и в эталоне, нет."""
    import docx
    p = load_project(YAML)
    p.raw["results"]["fill_series"] = [{"level_m": 7.2, "membrane": 100.0, "fiber": 150.0, "equivalent": 160.0},
                                       {"level_m": 11.0, "membrane": 200.0, "fiber": 250.0, "equivalent": 260.0},
                                       {"level_m": 17.1, "membrane": 900.0, "fiber": 950.0, "equivalent": 960.0}]
    path, _ = build_report(p, tmp_path / "f.docx")
    text = "\n".join(x.text for x in docx.Document(path).paragraphs)
    assert "5.3 Допустимая высота налива" not in text and "Допустимый уровень налива продукта по условию прочности составил" in text
    assert "Максимальные напряжения в стенке по уровням налива" not in text


def test_loads_cases_with_fill_levels():
    from ansys_report.loads.rvs_loads import compute
    from ansys_report.loads.export import build_cases
    from tests.test_calc import _points_project  # noqa: F401
    inp = {"diameter_m": 28.5, "wall_height_m": 18.0, "belt_height_m": 1.5, "n_belts": 12, "thicknesses_mm": [12] * 12, "product_density_t_m3": 0.934, "fill_height_m": 11.0,
           "hydrotest_height_m": 11.0, "max_fill_height_m": 11.0, "overpressure_pa": 0, "vacuum_pa": 0, "snow_district": "IV", "wind_district": "II", "terrain": "A", "aero_coef": 0.8,
           "roof_rise_m": 3.7, "roof_deck_thickness_m": 0.004, "ryn_mpa": 420, "e_mpa": 206000, "steel_density_kg_m3": 7850, "g": 9.8, "water_density_kg_m3": 1000,
           "gamma_m": 1.1, "gamma_n": 1.1, "gamma_c1": 0.7, "gamma_c2": 0.8, "gamma_c_rest": 0.8, "gamma_c_hydro": 0.9, "corrosion_weak_mm_year": 0.05,
           "corrosion_medium_mm_year": 0.1, "service_life_years": 40, "anticorrosion_life_years": 10, "diagnostics_interval_years": 10, "insulation_thickness_m": 0,
           "insulation_density_kg_m3": 140, "equipment": []}
    cs = build_cases(compute(inp), True, [7.2, 17.1])
    assert [c["name"] for c in cs][-2:] == ["Налив 7.2 м", "Налив 17.1 м"] and cs[-1]["hydro"]["free_surface_mm"] == 17100.0


def test_stability_ring_limit_matches_excel():
    import json
    from pathlib import Path
    from ansys_report.loads.rvs_loads import compute
    r = compute(json.loads((Path(__file__).parent.parent / "examples" / "loads_input_RVSP10000.json").read_text(encoding="utf-8")))
    b = r["stability"]["belts"]
    assert abs(b[0]["utilization"] - 1.2417) < 1e-3
    assert abs(b[5]["ring_hr_max_m"] - 0.5115) < 1e-3 and abs(b[7]["ring_hr_max_m"] - 2.369) < 1e-3
    assert abs(r["ring_modulus_cm3"] - 120.1) < 0.1


def test_report_with_analytic_stability(tmp_path):
    """Аналитическая проверка устойчивости (Excel ГОСТ 31385) в отчёт не выводится: в эталоне ТО-019-26 раздела 6.2 нет."""
    import json
    import docx
    from pathlib import Path
    from ansys_report.loads.rvs_loads import compute
    from ansys_report.loads.export import write_package
    res = compute(json.loads((Path(__file__).parent.parent / "examples" / "loads_input_RVSP10000.json").read_text(encoding="utf-8")))
    pk = write_package(res, tmp_path / "pk")
    ns: dict = {}
    exec(pk["data"].read_text(encoding="utf-8"), ns)
    p = load_project(YAML)
    p.raw["results"]["analytic"] = ns["LOADS"]["analytic"]
    path, _ = build_report(p, tmp_path / "a.docx")
    d = docx.Document(path)
    text = "\n".join(x.text for x in d.paragraphs)
    assert "6.2 Проверка устойчивости стенки по аналитической методике" not in text and len(d.tables) <= 9


def test_cad_fine_defect_large_tank_falls_back(tmp_path):
    """РВС-30000 с мелким шагом замеров на вмятине: CAD-сетка по шагу замеров слишком большая - должен сработать запасной вариант, а не ошибка."""
    from ansys_report.geometry import check
    pr = _points_project()
    pr["diameter"] = "45.6"
    xs = [round(-0.2 + 0.02 * i, 3) for i in range(21)]
    pr["defects"][0].update(nx="21", ny="21", x=xs, y=xs, values=[["-3" if 0 < i < 20 and 0 < j < 20 else "0" for i in range(21)] for j in range(21)])
    out = check.export_geometry(pr, tmp_path, 0.1, "STEP", lambda s: None)
    assert any(str(f).endswith(".step") for f in out["files"])


def test_agent_launch_is_windowless_and_survives_bad_link(tmp_path, monkeypatch):
    """Помощник под pythonw: нет консоли (stdout=None), дочерние процессы запускаются без окна, плохая ссылка не роняет запуск."""
    import subprocess
    import sys
    from ansys_report import agent
    monkeypatch.setattr(agent, "WORK", tmp_path)
    shown = []
    monkeypatch.setattr(agent, "popup", lambda title, text: shown.append(text))          # в тесте окно сообщения не показываем
    monkeypatch.setattr(sys, "stdout", None)
    assert agent.launch("ansysreport://run?x=1") == 1
    assert "Неверная ссылка" in (tmp_path / "agent.log").read_text(encoding="utf-8") and "Неверная ссылка" in shown[0]
    calls = []
    monkeypatch.setattr(subprocess, "call", lambda *a, **k: calls.append(k) or 0)
    bat = tmp_path / "x.bat"
    bat.write_text("@echo off")
    agent._run_bat(bat, {})
    assert calls[0]["creationflags"] == agent.NO_WINDOW and calls[0]["stdin"] == subprocess.DEVNULL


def test_pressure_view_matches_apply_loads_formula():
    """Показ нагрузок на сайте считает давление той же формулой, что скрипт наложения в Ansys: равномерное + гидростатика + ветер по поясам."""
    from ansys_report.loads.export import build_cases, pressure_view, report_fill_levels
    from ansys_report.loads.rvs_loads import compute
    res = compute(json.loads((Path(__file__).parent.parent / "examples/loads_input_RVSP10000.json").read_text(encoding="utf-8")))
    cases = build_cases(res, True, report_fill_levels(res))
    H = res["inputs"]["H"] * 1000
    for c in cases:
        v = pressure_view(c, H)
        hy = c.get("hydro")
        for z, p in zip(v["wall"]["z_m"], v["wall"]["p_kpa"]):
            zz = z * 1000
            want = c["uniform_mpa"] + (hy["gradient_mpa_mm"] * (hy["free_surface_mm"] - zz) if hy and zz < hy["free_surface_mm"] else 0) \
                + sum(w["p_mpa"] for w in c["wind"] if w["z_lo_mm"] <= zz < w["z_hi_mm"])
            assert abs(p - want * 1000) < 1e-3
        assert abs(v["roof_kpa"] - (c["uniform_mpa"] + c["roof_extra_mpa"]) * 1000) < 1e-3
        assert (v["bottom_kpa"] is None) == (hy is None)
    names = [pressure_view(c, H)["name"] for c in cases]
    assert names[0] == "Эксплуатация" and any(n.startswith("Налив") for n in names)


def test_pick_ansys_version(monkeypatch):
    import pytest
    from ansys_report.ansys_connect import ConnectError, pick_version
    inst = {261: "a", 242: "b"}
    monkeypatch.delenv("ANSYS_REPORT_VERSION", raising=False)
    assert pick_version(inst) == 261 and pick_version(inst, 242) == 242
    monkeypatch.setenv("ANSYS_REPORT_VERSION", "242")
    assert pick_version(inst) == 242
    monkeypatch.setenv("ANSYS_REPORT_VERSION", "251")
    with pytest.raises(ConnectError):
        pick_version(inst)


def test_check_ansys_journal_and_missing_install(monkeypatch, tmp_path):
    """check-ansys: журнал формируется без ошибок подстановки; без установленного Ansys - понятное сообщение."""
    import pytest
    from ansys_report import compat
    from ansys_report.ansys_connect import ConnectError
    from pathlib import PureWindowsPath
    j = compat.journal_text(PureWindowsPath(r"C:\ansys_report_work\out\wb_status.txt"), PureWindowsPath(r"C:\Ansys-report\ansys_scripts\probe_api.py"))
    compile(j, "probe.wbjn", "exec")                                   # журнал Workbench - корректный Python (\a, \x в путях Windows не портят строку)
    ns = {}
    cmd = next(l for l in j.splitlines() if "SendCommand" in l).split("Command=", 1)[1].rstrip(")")
    import re
    code = eval(cmd)                                                      # то, что выполнит Mechanical
    assert eval(re.search(r"io\.open\((.*), encoding", code).group(1)) == r"C:\Ansys-report\ansys_scripts\probe_api.py"      # путь доходит без искажений
    assert "probe_api.py" in j and compat.version_name(221) == "2022 R1" and compat.version_name(261) == "2026 R1"
    for k in list(__import__("os").environ):
        if k.startswith("AWP_ROOT"):
            monkeypatch.delenv(k)
    with pytest.raises(ConnectError, match="AWP_ROOT"):
        compat.check_ansys()
    compile(open(compat.SCRIPTS / "probe_api.py", encoding="utf-8").read(), "probe_api.py", "exec")


def test_load_figures_use_mpa(tmp_path):
    """Рисунки нагрузок в отчёте подписаны в МПа (как в эталонных отчётах), малые значения не округляются до нуля."""
    from ansys_report.loads.export import write_package
    from ansys_report.loads.figures import make_load_figures, mpa
    from ansys_report.loads.rvs_loads import compute
    assert mpa(0.168) == "0,168" and mpa(0.00238) == "0,00238" and mpa(0.000291) == "0,000291" and mpa(0) == "0" and mpa(-0.0017) == "0,00170"
    res = compute(json.loads((Path(__file__).parent.parent / "examples/loads_input_RVSP10000.json").read_text(encoding="utf-8")))
    pack = write_package(res, tmp_path / "pk", all_cases=True)
    figs = make_load_figures(pack["data"], tmp_path / "img")
    assert [f.name for f in figs] == ["fig05_hydrostatic.png", "fig04_loads.png"] and all(f.stat().st_size > 10000 for f in figs)
    assert "кПа" not in (Path(__file__).parent.parent / "ansys_report/loads/figures.py").read_text(encoding="utf-8")
