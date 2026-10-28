"""Структура Word-отчёта по эталону ТО-019-26 (Т-001): разделы, разрывы страниц, таблицы, формулы, данные заказчика."""
import json
from pathlib import Path

import docx
from docx.oxml.ns import qn

from ansys_report import report_input as RI
from ansys_report.project import load_project
from ansys_report.word_gen import build_report, g3

YAML = Path(__file__).parent.parent / "examples/demo/project.yaml"
T001 = {"number": "ТО-019-26", "year": "2026", "executor_name": "Иванов И. И.", "tank_name": "РВС-9985", "tag": "Т-001", "site": "Иркутский завод полимеров",
        "site_number": "5210-Т-001", "purpose": "хранения технологической воды", "product": "Техническая вода", "hazard_class": "КС-2а", "roof_radius": "34200",
        "steel_name": "09Г2С", "steel_standard": "ГОСТ 19281-2014", "steel_density": "7850", "steel_e_mpa": "206000", "steel_poisson": "0,3",
        "steel_ultimate_mpa": "480", "steel_yield_mpa": "325", "elongation_pct": "21", "steel_thickness_range": "9,01-25", "roof_mass_kg": "64156",
        "roof_equipment_kg": "819,5; 5151; 1571,8", "psi": "0,95", "roof_gf": "1,05", "load_factor": "1,3", "overp_gf": "1,2", "roof_ins_kg": "2080",
        "wall_ins_kg": "4335", "ins_thickness_mm": "100", "snow_kpa": "1,8", "snow_gf": "1,4", "snow_blowoff": "Не учитывается", "wind_kpa": "0,3"}
LOADS = {"overpressure_pa": 0, "terrain": "A", "product_density_t_m3": 1.0, "fill_height_m": 17.2}

# заголовки эталона по порядку (номер раздела выводов у нас 7: в эталоне пропущен 7 и стоит 8)
HEADINGS = ["1 Вводная часть", "2 Информация об объекте моделирования", "3 Сетка конструкции", "4 Граничные условия и нагрузки",
            "5 Результаты расчета на прочность", "5.1 Результаты расчета НДС РВС", "5.2 Результаты расчета НДС РВС", "6 Результаты расчета РВС на устойчивость",
            "6.1 Результаты расчета РВС на устойчивость", "7 Выводы и рекомендации", "Приложение 1", "Приложение 2", "Приложение 3"]
NEW_PAGE = ["3 Сетка", "4 Граничные", "5 Результаты", "5.2 Результаты", "6 Результаты", "7 Выводы", "Приложение 1", "Приложение 2", "Приложение 3"]


def _report(tmp_path, **over):
    data, err = RI.parse({**T001, **over})
    assert not err and not RI.missing(data)
    p = load_project(YAML)
    an = RI.anketa(data, LOADS)
    p.raw["report_data"] = json.loads(an["report_data"])
    p.tank.update(name=an["tank_name"], site=an["site"], site_number=an["site_number"], purpose=an["purpose"])      # как из project_template.yaml
    from ansys_report.project import _apply_report_steel
    _apply_report_steel(p.raw)
    path, _ = build_report(p, tmp_path / "ref.docx")
    return p, docx.Document(path)


def test_report_structure_matches_reference(tmp_path):
    p, d = _report(tmp_path)
    heads = [x for x in d.paragraphs if x.style.name.startswith("Heading")]
    texts = [x.text for x in heads]
    pos = [next(i for i, t in enumerate(texts) if t.startswith(h)) for h in HEADINGS]
    assert pos == sorted(pos)
    for h in NEW_PAGE:
        par = next(x for x in heads if x.text.startswith(h))
        assert par.paragraph_format.page_break_before, h
    body = "\n".join(x.text for x in d.paragraphs)
    # лишнее из прежней версии убрано
    for gone in ("Оценка отклонений образующих", "по уровням налива", "Коэффициенты запаса устойчивости", "аналитической методике", "Справочно", "Fixed Support"):
        assert gone not in body, gone
    caps = [x.text for x in d.paragraphs if x.text.startswith("Таблица ")]
    assert caps == ["Таблица 1 – Характеристика объекта исследования", "Таблица 2 – Толщина элементов",
                    "Таблица 3 – механические характеристики используемых материалов",
                    "Таблица 1.1 – Отклонения стенки от вертикали полного резервуара.", "Таблица 1.2 – Отклонения стенки от вертикали резервуара после опорожнения."]
    # подписи рисунков: сетка и результаты - с «РВС»
    assert "Рисунок 2. Сетка элементов стенки и днища РВС" in body and "Рисунок 6. Эквивалентные поверхностные напряжения в стенке РВС" in body


def test_report_tables_and_formulas_like_reference(tmp_path):
    p, d = _report(tmp_path)
    t1 = d.tables[1]                                       # [0] - подпись на титуле
    rows = {r.cells[0].text: r.cells[1].text for r in t1.rows}
    assert len(t1.rows) == 19
    assert rows["Радиус сферической крыши, мм"] == "34200" and rows["Наименование хранимого продукта"] == "Техническая вода"
    assert rows["Собственный фактический вес крыши, оборудования, площадок и ограждений на крыше, кг"] == "71698,3"
    assert rows["Тип местности"] == "А" and rows["Класс опасности резервуара по ГОСТ 31385-2023"] == "КС-2а" and rows["Расчетное ветровое давление, кПа"] == "0,3"
    t3 = d.tables[3]
    assert [c.text for c in t3.rows[1].cells] == ["09Г2С", "7850", "2,06х1011", "0,3", "480", "325", "9,01-25"]
    assert any(r.font.superscript for r in t3.rows[1].cells[2].paragraphs[0].runs)
    # таблицы по поясам: шапка в два уровня, номера столбцов, критерии [σ] и 3[σ]
    strength = [t for t in d.tables if t.rows[0].cells[0].text == "№ п/п"]
    assert len(strength) == 3
    fib, eq, mem = strength
    assert "фибровые" in fib.rows[0].cells[2].text and fib.rows[1].cells[3].text.endswith("3[σ]") and [c.text for c in fib.rows[2].cells] == ["1", "2", "3", "4", "5"]
    assert mem.rows[1].cells[3].text.endswith(", [σ]") and mem.rows[3].cells[3].text == "216,7" and fib.rows[3].cells[3].text == "650"
    assert len(mem.rows) == 3 + len(p.tank["belts_thickness"])
    # формулы редактора Word: R, E2, Fсн, F
    m = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
    formulas = ["".join(t.text or "" for t in om.iter(m + "t")) for om in d.element.body.iter(m + "oMath")]
    assert any(f.startswith("R=325∙0.7∙1.01.05∙1.05=206,35") for f in formulas)
    assert any(f.startswith("E2=σвр-σу∆100-σуE=") for f in formulas)
    assert any(f.startswith("Fсн = r∙S2 = 14,25∙21422 = 15262 Н/м") for f in formulas)
    roof = next(f for f in formulas if f.startswith("F = γf"))
    assert "1,05∙(64156+0,95∙(819,5+5151+1571,8))∙9,81 = 734643,9 Н." in roof
    body = "\n".join(x.text for x in d.paragraphs)
    assert "Резервуар РВС-9985 установлен на объекте: «Иркутский завод полимеров» (№ 5210-Т-001)" in body
    assert "Стенка резервуара состоит из листов стали 09Г2С, ГОСТ 19281-2014." in body
    assert "Fизол ст = 4335∙9,81 = 42526 Н" in body


def test_g3_numbers_like_reference():
    assert [g3(x) for x in (7.45, 14.9, 89.4, 0.95, 1.05, 64156, 1571.8, 0.3, 216.66)] == ["7,45", "14,9", "89,4", "0,95", "1,05", "64156", "1571,8", "0,3", "216,7"]


def test_reference_figures_folder_has_priority(tmp_path):
    """Рисунки export_figures.py (figures/fig01.png …, устойчивость fig14…) берутся в отчёт вместо старых снимков images/."""
    import shutil
    from PIL import Image
    root = tmp_path / "proj"
    shutil.copytree(YAML.parent, root)
    (root / "figures").mkdir(exist_ok=True)
    for name in ("fig01", "fig04", "fig14"):
        Image.new("RGB", (40, 30), (10, 120, 10)).save(root / "figures" / f"{name}.png")
    from ansys_report.word_gen import Report
    rep = Report(load_project(root / "project.yaml"))
    assert rep._find_image("fig01_general") == root / "figures" / "fig01.png"
    assert rep._find_image("fig04_loads") != root / "figures" / "fig04.png"          # рис. 4-5 - наши схемы нагрузок
    assert rep._find_image("fig_stab1") == root / "figures" / "fig14.png"
    assert rep._find_image("fig05_hydrostatic") != root / "figures" / "fig05.png"


def test_snow_line_load_like_reference(tmp_path):
    """Снег как в эталоне: погонная сила Fсн = 11 610 Н/м на кромку - в расчёт нагрузок (S = 2·Fсн/r) и в раздел 4."""
    import math
    from ansys_report.loads.rvs_loads import compute
    from tests.test_calc import _points_project  # noqa: F401
    inp = {"diameter_m": 28.5, "wall_height_m": 17.88, "belt_height_m": 1.49, "n_belts": 12, "thicknesses_mm": [16, 14, 13, 12, 12, 12, 12, 12, 11, 11, 11, 11],
           "product_density_t_m3": 1.0, "fill_height_m": 17.2, "overpressure_pa": 0, "vacuum_pa": 0, "snow_district": "IV", "wind_district": "II",
           "terrain": "A", "roof_rise_m": 3.11, "roof_deck_thickness_m": 0.004, "ryn_mpa": 325, "snow_kpa": 1.8, "snow_line_n_m": 11610}
    L = compute(inp)["loads"]
    r = 14.25
    total = L["snow_mpa"]["design"] * 1e6 * math.pi * r * r                  # что ляжет на кромку (S·πR²)
    assert abs(total - 11610 * 2 * math.pi * r) < 1
    _, d = _report(tmp_path, snow_line_n_m="11610")
    m = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
    formulas = ["".join(t.text or "" for t in om.iter(m + "t")) for om in d.element.body.iter(m + "oMath")]
    assert any(f.startswith("Fсн = r∙S2 = 14,25∙16292") and f.endswith("= 11610 Н/м.") for f in formulas)
