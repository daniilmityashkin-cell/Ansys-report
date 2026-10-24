from pathlib import Path
"""Сайт: регистрация, замеры → проверка → оплата (тестовый провайдер) → геометрия, доступ только владельцу, пакеты генераций."""
import json
import re

import pytest
from fastapi.testclient import TestClient

from server.app import create_app
from server.config import Settings
from tests.test_calc import _points_project


@pytest.fixture()
def site(tmp_path):
    app = create_app(Settings(data_dir=tmp_path, inline_jobs=True, secret_key="t" * 32))
    return app


def csrf_of(c, path):
    html = c.get(path).text
    return re.search(r'name="csrf" value="([^"]+)"', html).group(1)


def register(c, email="a@example.com", password="correct horse battery"):
    tok = csrf_of(c, "/register")
    return c.post("/register", data={"email": email, "password": password, "consent": "1", "csrf": tok}, follow_redirects=False)


def test_register_login_logout_and_csrf(site):
    c = TestClient(site)
    assert c.post("/register", data={"email": "x@example.com", "password": "correct horse battery", "consent": "1"}).status_code == 400   # без CSRF
    r = register(c)
    assert r.status_code == 303 and r.headers["location"] == "/cabinet"
    assert "Личный кабинет" in c.get("/cabinet").text
    assert register(TestClient(site)).status_code == 400                                       # почта уже занята
    c2 = TestClient(site)
    tok = csrf_of(c2, "/register")
    assert c2.post("/register", data={"email": "b@example.com", "password": "short", "consent": "1", "csrf": tok}).status_code == 400
    assert c2.post("/register", data={"email": "b@example.com", "password": "long enough pass", "csrf": tok}).status_code == 400    # нет согласия
    tok = csrf_of(c, "/cabinet")
    assert c.post("/logout", data={"csrf": tok}, follow_redirects=False).status_code == 303
    assert c.get("/cabinet", follow_redirects=False).status_code == 303
    c3 = TestClient(site)
    tok = csrf_of(c3, "/login")
    bad = c3.post("/login", data={"email": "a@example.com", "password": "wrong password!", "csrf": tok})
    assert bad.status_code == 400
    ok = c3.post("/login", data={"email": "a@example.com", "password": "correct horse battery", "csrf": tok}, follow_redirects=False)
    assert ok.status_code == 303


def new_generation(c, project=None, title="РВС-1"):
    tok = csrf_of(c, "/new")
    body = json.dumps(project or _points_project(), ensure_ascii=False)
    return c.post("/new", data={"title": title, "text": body, "csrf": tok}, follow_redirects=False)


def test_full_flow_pay_and_download(site):
    c = TestClient(site)
    register(c)
    bad = c.post("/new", data={"title": "x", "text": "не json", "csrf": csrf_of(c, "/new")})
    assert bad.status_code == 400 and "Не удалось прочитать" in bad.text
    r = new_generation(c)
    assert r.status_code == 303
    gen = r.headers["location"]
    page = c.get(gen).text
    assert "Результат проверки" in page and "Получить отчёт и файлы" in page
    assert c.get(gen + "/all.zip").status_code == 404                              # до оплаты файлов нет
    tok = csrf_of(c, gen)
    pay = c.post(gen + "/pay", data={"csrf": tok}, follow_redirects=False)
    assert pay.status_code == 303 and pay.headers["location"].startswith("/pay/fake/")
    pay_url = pay.headers["location"]
    assert "Тестовая оплата" in c.get(pay_url).text
    done = c.post(pay_url, data={"csrf": csrf_of(c, pay_url)}, follow_redirects=False)
    assert done.status_code == 303 and done.headers["location"] == gen
    page = c.get(gen).text
    assert "готово" in page and "Поверхность РВС.step" in page
    zr = c.get(gen + "/all.zip")
    assert zr.status_code == 200 and zr.content[:2] == b"PK"
    again = c.post(pay_url, data={"csrf": csrf_of(c, pay_url)}, follow_redirects=False)       # повтор оплаты не запускает расчёт заново
    assert again.status_code == 303
    for evil in ("/files/..%2F..%2Finput.json", "/files/%2e%2e/%2e%2e/input.json", "/files/../../input.json"):
        assert c.get(gen + evil).status_code in (404, 400, 422)                                # выход из папки результата невозможен


def test_other_user_cannot_see_generation(site):
    a = TestClient(site)
    register(a, "a@example.com")
    gen = new_generation(a).headers["location"]
    b = TestClient(site)
    register(b, "b@example.com")
    assert b.get(gen).status_code == 404
    assert b.get(gen + "/all.zip").status_code == 404
    anon = TestClient(site)
    assert anon.get(gen, follow_redirects=False).status_code == 303


def _buy_pack(c, size=10):
    tok = csrf_of(c, "/packs")
    pay = c.post("/packs/buy", data={"csrf": tok, "size": str(size)}, follow_redirects=False).headers["location"]
    c.post(pay, data={"csrf": csrf_of(c, pay)})
    return pay


def test_pack_purchase_and_credits(site):
    c = TestClient(site)
    register(c)
    gen = new_generation(c).headers["location"]
    assert "входит в оплаченную генерацию" in c.get(gen + "/report").text
    pay_url = _buy_pack(c, 10)
    assert "Осталось генераций: <b>10</b>" in c.get("/packs").text
    r = c.post(gen + "/pay", data={"csrf": csrf_of(c, gen)}, follow_redirects=False)    # из пакета, без платежа
    assert r.headers["location"] == gen
    with site.state.Session() as db:
        from server.models import User, Payment
        assert db.query(User).one().credits == 9
        assert db.query(Payment).count() == 1
    c.post(pay_url, data={"csrf": csrf_of(c, pay_url)})            # тот же платёж повторно: пакет не зачисляется дважды
    with site.state.Session() as db:
        assert db.query(User).one().credits == 9
    bad = c.post("/packs/buy", data={"csrf": csrf_of(c, "/packs"), "size": "7"}, follow_redirects=False)
    assert bad.headers["location"] == "/packs"                       # неизвестный размер пакета не продаётся
    with site.state.Session() as db:
        assert db.query(Payment).count() == 1                         # новый платёж не создан


def test_table_input_to_generation(site):
    from server.tableinput import build_project
    # 5 образующих × 3 пояса, строки - образующие (как у заказчика): каждая строка - одна образующая по поясам
    table = "-10\t-9\t-6\n-12\t-23\t-27\n1\t-5\t-20\n60\t60\t52\n53\t64\t73"
    pr = build_project("28,5", "18", table, "13 12 11", "gens")
    assert len(pr["angles"]) == 5 and len(pr["heights"]) == 4 and pr["values"][0] == ["0"] * 5 and len(pr["bands"]) == 3
    assert pr["values"][1][0] == "-10.0" and pr["values"][3][4] == "73.0"        # столбец-пояс 3 у образующей 5
    for bad in (("0", "18", table, "13 12 11"), ("28.5", "18", table, "13 12"), ("28.5", "18", "1 2\n3", "13"), ("28.5", "18", "", "13")):
        with pytest.raises(ValueError):
            build_project(*bad)
    c = TestClient(site)
    register(c)
    tok = csrf_of(c, "/new/table")
    r = c.post("/new/table", data={"title": "РВС-2", "diameter": "28.5", "height": "18", "table": table, "thicknesses": "13 12 11",
                                   "rows_are": "gens", "csrf": tok}, follow_redirects=False)
    assert r.status_code == 303
    assert "Результат проверки" in c.get(r.headers["location"]).text
    bad = c.post("/new/table", data={"title": "x", "diameter": "28.5", "height": "18", "table": table, "thicknesses": "13",
                                     "rows_are": "gens", "csrf": tok})
    assert bad.status_code == 400 and "Толщин указано" in bad.text


def test_reconcile_applies_missed_payment(site, monkeypatch):
    """Уведомление об оплате не дошло: сверка с провайдером всё равно запускает генерацию (платёж не теряется)."""
    import server.app as appmod
    from server.models import Payment, Generation

    class Stub:
        def verify(self, provider_id):
            return "paid"
    monkeypatch.setattr(appmod, "get_provider", lambda name: Stub())
    c = TestClient(site)
    register(c)
    gen = new_generation(c).headers["location"]
    gid = int(gen.rsplit("/", 1)[1])
    with site.state.Session() as db:
        u = db.get(Generation, gid).user_id
        db.add(Payment(user_id=u, kind="generation", generation_id=gid, amount_rub=1500, provider="stub", provider_id="X1"))
        db.commit()
    assert site.state.reconcile() == 1
    assert site.state.reconcile() == 0                                   # повторно ничего не применяется
    assert "готово" in c.get(gen).text


def test_table_from_excel(site):
    """Замеры из файла Excel: заголовки и номера отбрасываются, остаётся блок чисел."""
    import io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Таблица 1.1 Отклонения, мм"])
    ws.append(["№ обр.", 1, 2, 3])
    for g, row in enumerate([(-10, -9, -6), (-12, -23, -27), (1, -5, -20), (60, 60, 52), ("53,5", 64, 73)], 1):
        ws.append([g, *row])
    buf = io.BytesIO()
    wb.save(buf)
    c = TestClient(site)
    register(c)
    tok = csrf_of(c, "/new/table")
    r = c.post("/new/table", data={"csrf": tok, "title": "XLS", "diameter": "28.5", "height": "18", "thicknesses": "13 12 11", "rows_are": "gens"},
               files={"xlsx": ("замеры.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, follow_redirects=False)
    assert r.status_code == 303, r.text[:300]
    page = c.get(r.headers["location"]).text
    assert "15 (5 образующих × 3 поясов)" in page
    bad = c.post("/new/table", data={"csrf": tok, "diameter": "28.5", "height": "18", "thicknesses": "13 12 11"},
                 files={"xlsx": ("x.xlsx", b"not excel", "application/octet-stream")})
    assert bad.status_code == 400 and "Excel" in bad.text


def _pay_fake(c, gen):
    r = c.post(gen + "/pay", data={"csrf": csrf_of(c, gen)}, follow_redirects=False)
    if r.headers["location"].startswith("/pay/fake/"):
        c.post(r.headers["location"], data={"csrf": csrf_of(c, r.headers["location"])})
    return r


def test_one_generation_per_account_then_pack(site):
    c = TestClient(site)
    register(c)
    g1 = new_generation(c).headers["location"]
    _pay_fake(c, g1)
    assert "готово" in c.get(g1).text
    p2 = _points_project()
    p2["diameter"] = "31.0"                                    # другие данные: проверяем именно правило «одна на аккаунт»
    g2 = new_generation(c, p2, "РВС-2").headers["location"]
    page = c.get(g2).text
    assert "уже использована" in page and "1\u00a0500" not in page
    r = c.post(g2 + "/pay", data={"csrf": csrf_of(c, g2)}, follow_redirects=False)
    assert r.headers["location"] == g2                         # платёж не создаётся
    with site.state.Session() as db:
        from server.models import Payment
        assert db.query(Payment).count() == 1
    _buy_pack(c, 10)                                            # пакет снимает ограничение
    assert "из пакета" in c.get(g2).text


def test_same_data_on_other_account_needs_pack(site):
    a, b = TestClient(site), TestClient(site)
    register(a)
    _pay_fake(a, new_generation(a).headers["location"])
    register(b, "b@example.com")
    g = new_generation(b).headers["location"]                  # те же замеры на другом аккаунте
    assert "другом аккаунте" in b.get(g).text


def test_registration_abuse_protection(site):
    c = TestClient(site)
    assert register(c, "a.b+one@gmail.com").status_code == 303
    assert register(TestClient(site), "ab@gmail.com").status_code == 400               # тот же адрес с другой записью
    assert register(TestClient(site), "x@mailinator.com").status_code == 400            # одноразовая почта
    assert register(TestClient(site), "c@example.com").status_code == 303
    assert register(TestClient(site), "d@example.com").status_code == 303
    assert register(TestClient(site), "e@example.com").status_code == 400               # 4-й аккаунт с одного адреса за сутки


def test_old_database_gets_new_columns(tmp_path):
    import sqlite3
    db = tmp_path / "site.db"
    con = sqlite3.connect(db)
    con.execute("CREATE TABLE users (id INTEGER PRIMARY KEY, email VARCHAR(254), password_hash VARCHAR(255), created DATETIME, subscription_until DATETIME, consent_at DATETIME)")
    con.execute("INSERT INTO users (email, password_hash) VALUES ('old@example.com', 'x')")
    con.commit(); con.close()
    app = create_app(Settings(data_dir=tmp_path, inline_jobs=True, secret_key="t" * 32))
    with app.state.Session() as s:
        from server.models import User
        assert s.query(User).one().email_key == "old@example.com"


def test_preview_free_before_payment(site):
    c = TestClient(site)
    register(c)
    g = new_generation(c).headers["location"]
    page = c.get(g).text
    assert "3D-модель по результатам съёмки" in page and "Получить отчёт и файлы" in page
    img = c.get(g + "/img/prev_unfold.png")
    assert img.status_code == 200 and img.headers["content-type"] == "image/png"
    assert c.get(g + "/all.zip").status_code == 404                         # файлы - только после оплаты
    assert TestClient(site).get(g + "/img/prev_unfold.png", follow_redirects=False).status_code in (303, 404)


def test_viewer_json_owner_only(site):
    c = TestClient(site)
    register(c)
    g = new_generation(c).headers["location"]
    r = c.get(g + "/viewer.json")
    assert r.status_code == 200
    d = r.json()
    assert d["nt"] and d["nz"] and len(d["grid"]) == d["nz"] * d["nt"]
    gid = g.rsplit("/", 1)[1]
    vj = next(Path(str(x)) for x in site.state.settings.data_dir.glob(f"generations/{gid}/preview/viewer.json"))
    vj.unlink()                                                              # старая генерация без 3D-данных
    assert c.get(g + "/viewer.json").status_code == 200 and vj.exists()
    assert c.get(g + "/img/dent1.png").status_code in (200, 404) and c.get(g + "/img/../x.png").status_code == 404
    assert c.get("/static/vendor/three.module.min.js").status_code == 200
    assert TestClient(site).get(g + "/viewer.json", follow_redirects=False).status_code in (303, 404)


def test_loads_in_table_form_and_package(site):
    """Нагрузки: вводятся в форме, считаются на странице бесплатно, пакет для Ansys выдаётся после оплаты."""
    c = TestClient(site)
    register(c)
    tok = csrf_of(c, "/new/table")
    data = {"csrf": tok, "title": "РВС-Н", "diameter": "28.5", "height": "18", "thicknesses": "13 12 11", "rows_are": "gens",
            "table": "-10\t-9\t-6\n-12\t-23\t-27\n1\t-5\t-20\n60\t60\t52\n53\t64\t73",
            "l_on": "1", "l_density": "0.934", "l_fill": "16", "l_hydro": "16", "l_overp": "2000", "l_vac": "250", "l_snow": "III", "l_wind": "II",
            "l_terrain": "A", "l_roof_rise": "3", "l_deck": "4", "l_ryn": "345", "l_ins_t": "0", "l_ins_rho": "140",
            "l_equipment": "Люк-лаз 600; 2; 170; 1\nПлощадка на крыше; 1; 4297; крыша"}
    r = c.post("/new/table", data=data, follow_redirects=False)
    assert r.status_code == 303, r.text[:400]
    gen = r.headers["location"]
    page = c.get(gen).text
    assert "Нагрузки на резервуар" in page and "район III" in page and "Нагрузки по поясам стенки" in page and "Гидростатика, гидроиспытания, кПа" in page
    _pay_fake(c, gen)
    files = c.get(gen).text
    assert "loads/loads_data.py" in files and "loads/apply_loads.py" in files and "Как_наложить_нагрузки.txt" in files
    body = c.get(gen + "/files/loads/loads_data.py").text
    assert "LOADS = " in body
    bad = dict(data, l_snow="XX", csrf=csrf_of(c, "/new/table"))
    assert c.post("/new/table", data=bad).status_code == 400
    bad = dict(data, l_equipment="только название", csrf=csrf_of(c, "/new/table"))
    assert "Оборудование" in c.post("/new/table", data=bad).text


def test_open_in_ansys_package_in_paid_files(site):
    """В платных файлах - папка «Открыть в Ansys»: модель, толщины, нагрузки, скрипты и .bat."""
    c = TestClient(site)
    register(c)
    tok = csrf_of(c, "/new/table")
    data = {"csrf": tok, "title": "РВС-А", "diameter": "28.5", "height": "18", "thicknesses": "13 12 11", "rows_are": "gens",
            "table": "-10\t-9\t-6\n-12\t-23\t-27\n1\t-5\t-20\n60\t60\t52\n53\t64\t73", "l_on": "1", "l_density": "0.934", "l_fill": "16", "l_hydro": "16",
            "l_overp": "0", "l_vac": "0", "l_snow": "III", "l_wind": "II", "l_terrain": "A", "l_roof_rise": "3", "l_deck": "4", "l_ryn": "345",
            "l_ins_t": "0", "l_ins_rho": "140"}
    gen = c.post("/new/table", data=data, follow_redirects=False).headers["location"]
    _pay_fake(c, gen)
    page = c.get(gen).text
    for name in ("ansys_model/Make_report.bat", "ansys_model/Run_all.bat", "ansys_model/anketa.json", "ansys_model/Open_in_Ansys.bat", "ansys_model/model.step", "ansys_model/thickness_zones.csv", "ansys_model/loads_data.py",
                 "ansys_model/build_model.wbjn", "ansys_model/build_model.py", "ansys_model/apply_loads.py", "ansys_model/README_RU.txt"):
        assert name in page, name
    bat = c.get(gen + "/files/ansys_model/Open_in_Ansys.bat").text
    assert "RunWB2.exe" in bat and "AWP_ROOT261" in bat and "ANSYS_REPORT_AXIS=Z" in bat
    import ast
    ast.parse(c.get(gen + "/files/ansys_model/build_model.wbjn").text)
    assert "\n" in c.get(gen + "/files/ansys_model/Open_in_Ansys.bat").text


def _survey_book(names=("Т-009 пустой",)):
    """Книга в формате геодезической съёмки: блоки по поясам, пустая строка между блоками, замыкание (точка 1 в конце), округление высоты 2981."""
    import io
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for name in names:
        ws = wb.create_sheet(name)
        ws["B1"], ws["C1"] = 14250, 28500
        row = 3
        for belt, level in enumerate((1490, 2980, 4470), 1):
            for k in range(1, 16):
                ang = 85 - 24 * (k - 1)
                ws.cell(row, 2, k); ws.cell(row, 3, -10 + belt * k % 7 * 3); ws.cell(row, 4, ang); ws.cell(row, 5, 14240)
                ws.cell(row, 7, 2981 if (belt == 2 and k == 15) else level)
                row += 1
            ws.cell(row, 2, 1); ws.cell(row, 3, -10); ws.cell(row, 4, -275); ws.cell(row, 7, level)      # замыкание
            row += 2
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_survey_workbook_upload(site):
    c = TestClient(site)
    register(c)
    xl = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    base = {"csrf": csrf_of(c, "/new/table"), "title": "Т-009", "thicknesses": "13 12 11", "rows_are": "belts"}
    r = c.post("/new/table", data=base, files={"xlsx": ("съёмка.xlsx", _survey_book(), xl)}, follow_redirects=False)
    assert r.status_code == 303, r.text[:300]
    page = c.get(r.headers["location"]).text
    assert "45 (15 образующих × 3 поясов)" in page and "28.5 × 4.47 м" in page
    two = _survey_book(("Т-001 полный", "Т-001 пустой"))
    bad = c.post("/new/table", data=dict(base, csrf=csrf_of(c, "/new/table")), files={"xlsx": ("s.xlsx", two, xl)})
    assert bad.status_code == 400 and "несколько листов" in bad.text
    ok = c.post("/new/table", data=dict(base, csrf=csrf_of(c, "/new/table"), sheet="Т-001 пустой"), files={"xlsx": ("s.xlsx", two, xl)}, follow_redirects=False)
    assert ok.status_code == 303
    from server.tableinput import survey_from_xlsx
    assert survey_from_xlsx(two, "Т-001")["sheet"] == "Т-001 пустой"                 # без состояния - «пустой»
    assert survey_from_xlsx(two, "t-001 полный")["sheet"] == "Т-001 полный"           # латинская t, регистр
    with pytest.raises(ValueError):
        survey_from_xlsx(two, "Т-777")


def test_generation_status_endpoint(site):
    c = TestClient(site)
    register(c)
    g = new_generation(c).headers["location"]
    st = c.get(g + "/status").json()
    assert st["status"] == "draft" and st["preview"] is True and st["preview_error"] == ""
    assert TestClient(site).get(g + "/status", follow_redirects=False).status_code in (303, 404)
    other = TestClient(site)
    register(other, "b@example.com")
    assert other.get(g + "/status").status_code == 404                                # чужая генерация
    _pay_fake(c, g)
    assert c.get(g + "/status").json()["status"] == "done"


def test_structure_form_and_json_project(site):
    """Конструкция вводится формой или берётся из файла проекта; 3D-схема и таблицы показываются на странице."""
    c = TestClient(site)
    register(c)
    data = {"csrf": csrf_of(c, "/new/table"), "title": "К", "diameter": "28.5", "height": "18", "thicknesses": "13 12 11", "rows_are": "gens", "src": "manual",
            "table": "-10\t-9\t-6\n-12\t-23\t-27\n1\t-5\t-20\n60\t60\t52\n53\t64\t73",
            "s_nz_name": ["П1"], "s_nz_shape": ["Круглый"], "s_nz_width": ["426"], "s_nz_height": [""], "s_nz_z": ["1.1"], "s_nz_angle": ["353"], "s_nz_t": ["11"],
            "s_nz_in": ["200"], "s_nz_out": ["500"], "s_posted": "1", "s_bot_on": "1", "s_bot_outside": "200", "s_bot_width": "1500", "s_bot_overlap": "100",
            "s_bot_slope": "100", "s_bot_plate_t": "8", "s_bot_ann_t": "12", "s_roof_on": "1", "s_roof_shape": "Сферическая", "s_roof_sphere_r": "50",
            "s_roof_cone_rise": "2", "s_roof_central_r": "1.3", "s_roof_radial": "24", "s_roof_first": "7.5", "s_roof_rings": "10", "s_roof_deck_t": "7",
            "s_roof_band_w": "1000", "s_roof_band_t": "10", "s_roof_band_drop": "200", "s_roof_beam_drop": "20", "s_roof_deck_inset": "200",
            "s_beam_radial": "Двутавр 25Б2", "s_beam_central": "Швеллер 24П", "s_beam_rings": "Уголок L90×90×6", "s_beam_edge": "Швеллер 24П"}
    r = c.post("/new/table", data=data, follow_redirects=False)
    assert r.status_code == 303, r.text[:300]
    g = r.headers["location"]
    page = c.get(g).text
    assert "Конструкция резервуара" in page and "Ø426" in page and "Сечения балок крыши" in page
    js = c.get(g + "/structure.json").json()
    assert js["nozzles"][0]["angle"] == 353 and js["roof"]["rise_m"] > 2
    assert TestClient(site).get(g + "/structure.json", follow_redirects=False).status_code in (303, 404)
    bad = dict(data, csrf=csrf_of(c, "/new/table"), s_nz_z=["99"])
    assert "не помещается" in c.post("/new/table", data=bad).text
    proj = json.dumps(_points_project(), ensure_ascii=False).encode()
    r2 = c.post("/new/table", data={"csrf": csrf_of(c, "/new/table"), "src": "json", "title": "J"}, files={"jsonfile": ("p.json", proj, "application/json")}, follow_redirects=False)
    assert r2.status_code == 303
    assert c.post("/new/table", data={"csrf": csrf_of(c, "/new/table"), "src": "json"}).status_code == 400


def test_unlimited_account(tmp_path):
    """Аккаунт из ASR_UNLIMITED_EMAILS: любое число генераций без оплаты и пакетов."""
    app = create_app(Settings(data_dir=tmp_path, inline_jobs=True, secret_key="t" * 32, unlimited_emails=frozenset({"vip@example.com"})))
    c = TestClient(app)
    register(c, email="vip@example.com")
    for _ in range(3):
        g = new_generation(c).headers["location"]
        assert "безлимитный аккаунт" in c.get(g).text
        r = c.post(g + "/pay", data={"csrf": csrf_of(c, g)}, follow_redirects=False)
        assert r.status_code == 303
        assert "Скачать всё архивом" in c.get(g).text
    other = TestClient(app)
    register(other, email="x@example.com")
    g = new_generation(other).headers["location"]
    assert "безлимитный" not in other.get(g).text


def test_agent_flow(site):
    """Кнопка «Собрать отчёт в Ansys»: ссылка на помощника, скачивание пакета по токену, статус и загрузка отчёта обратно."""
    import re as _re
    c = TestClient(site)
    register(c)
    gen = new_generation(c).headers["location"]
    pay = c.post(gen + "/pay", data={"csrf": csrf_of(c, gen)}, follow_redirects=False).headers["location"]
    c.post(pay, data={"csrf": csrf_of(c, pay)}, follow_redirects=False)
    page = c.get(gen).text
    m = _re.search(r'href="(ansysreport://run\?server=([^&]+)&amp;g=(\d+)&amp;t=([0-9a-f]+))"', page)
    assert m, "на странице нет ссылки помощника"
    gid, tok = int(m.group(3)), m.group(4)
    assert c.get(f"/agent/{gid}/package.zip?t=bad").status_code == 404                    # без верного токена пакет не отдаётся
    assert c.post(f"/agent/{gid}/status?t=bad", json={"stage": "x"}).status_code == 403
    assert c.post(f"/agent/{gid}/status?t={tok}", json={"stage": "solve", "message": "считаю"}).json() == {"ok": True}
    up = c.post(f"/agent/{gid}/upload?t={tok}", files={"file": ("Report.docx", b"docx-bytes")})
    assert up.status_code == 200
    assert c.post(f"/agent/{gid}/upload?t={tok}", files={"file": ("evil.exe", b"x")}).status_code == 400
    st = c.get(f"/g/{gid}/agent.json").json()
    assert st["stage"] == "solve" and st["files"] == ["Report.docx"]
    assert c.get(f"/g/{gid}/agent-files/Report.docx").content == b"docx-bytes"
    assert TestClient(site).get(f"/g/{gid}/agent-files/Report.docx").status_code in (303, 404)   # чужой не скачает
    assert "Помощник для Ansys" in c.get("/agent/setup").text


def test_agent_link_parse():
    from ansys_report.agent import parse_link
    p = parse_link("ansysreport://run?server=https://x.example/&g=5&t=abc")
    assert p == {"server": "https://x.example", "gid": 5, "token": "abc"}
    with pytest.raises(ValueError):
        parse_link("http://evil/?server=1&g=1&t=1")
