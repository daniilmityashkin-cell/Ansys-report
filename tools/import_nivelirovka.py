"""Импорт замеров из Excel-образца «Нивелировка Т001-Т004.xlsx» в проекты examples/T00N_ansys.
Запуск: python tools/import_nivelirovka.py "Нивелировка Т001-Т004.xlsx"
Для каждого листа «Т-00N полный/пустой» берутся значения отклонений (мм) по поясам и точкам."""
import csv, re, shutil, sys
from pathlib import Path
import openpyxl

src = Path(sys.argv[1])
wb = openpyxl.load_workbook(src, data_only=True)
root = Path(__file__).resolve().parent.parent / "examples"
base = (root / "T001_ansys" / "project.yaml").read_text(encoding="utf-8")


def read_sheet(ws):
    marks = [r for r in range(1, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, (int, float)) and ws.cell(r, 1).value >= 1]
    stride = marks[2] - marks[1]
    n = stride - 2                      # точки без замыкающей и пустой строки
    belts = []
    for k, m in enumerate(marks):
        first = m + 1 if k == 0 else m  # у 1-го пояса данные со следующей строки
        belts.append([ws.cell(first + i, 3).value for i in range(n)])
    return n, belts


def angle_step(ws, n):
    return 360 / n


for tag in ("Т-001", "Т-002", "Т-003", "Т-004"):
    out = root / f"{tag.replace('-', '')}_ansys".replace("Т", "T")
    out.mkdir(exist_ok=True)
    ns = set()
    for state, fname in (("полный", "survey_full.csv"), ("пустой", "survey_empty.csv")):
        n, belts = read_sheet(wb[f"{tag} {state}"])
        ns.add(n)
        with (out / fname).open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["point"] + [f"belt{b}" for b in range(1, len(belts) + 1)])
            for i in range(n):
                w.writerow([i + 1] + [belts[b][i] for b in range(len(belts))])
    assert len(ns) == 1, f"{tag}: разное число точек {ns}"
    n = ns.pop()
    y = base.replace("Т-001", tag).replace("№Т-001", f"№{tag}").replace("points_per_belt: 15", f"points_per_belt: {n}")
    y = re.sub(r"(  tag: )\"?Т-001\"?", rf'\1"{tag}"', y)
    (out / "project.yaml").write_text(y, encoding="utf-8")
    print(tag, "точек на пояс:", n, "->", out)
