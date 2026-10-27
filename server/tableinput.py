"""Ввод замеров заказчиком на сайте: таблица из Excel + размеры → проект замеров (тот же формат, что и JSON)."""
from __future__ import annotations
import io
import re

from ansys_report.geometry.rvs_core import number

MAX_GENS, MAX_BELTS = 160, 100


def _cells(text: str) -> list[list[str]]:
    rows = []
    for line in (text or "").replace("\r", "").split("\n"):
        if not line.strip():
            continue
        parts = re.split(r"\t|;", line) if ("\t" in line or ";" in line) else re.split(r"\s+", line.strip())
        rows.append([p.strip() for p in parts])
    return rows


def parse_thicknesses(text: str) -> list[float]:
    vals = [v for v in re.split(r"[\s;]+", (text or "").strip()) if v]
    if not vals:
        raise ValueError("Укажите толщины поясов, мм (через пробел или запятую с пробелом: 13 12 12 11 …).")
    out = []
    for i, v in enumerate(vals, 1):
        t = number(v.rstrip(","), f"Толщина пояса {i}")
        if t <= 0:
            raise ValueError(f"Толщина пояса {i} должна быть положительной.")
        out.append(t)
    return out


def require_thicknesses(project: dict) -> None:
    """Толщины поясов обязательны: молча подставленные значения дали бы неверные напряжения в отчёте."""
    miss = []
    for i, b in enumerate(project.get("bands") or [], 1):
        try:
            ok = float(str(b[1]).replace(",", ".").strip()) > 0
        except (IndexError, ValueError, TypeError):
            ok = False
        if not ok:
            miss.append(i)
    if miss:
        raise ValueError("Не заданы толщины поясов: " + ", ".join(map(str, miss)) + ". Введите толщину каждого пояса, мм (колонка «t, мм» на шаге 01): по паспорту резервуара или толщинометрии.")


def build_project(diameter: str, height: str, table: str, thicknesses: str, rows_are: str = "gens") -> dict:
    """rows_are: 'gens' - строки таблицы = образующие, столбцы = пояса (как в таблицах заказчика);
    'belts' - строки = пояса, столбцы = образующие. Отклонения в мм: + наружу, − внутрь.
    Образующие равномерно по окружности (0°, 360°/n, …), пояса равной высоты H/число поясов."""
    D = number(diameter, "Диаметр")
    H = number(height, "Высота")
    if D <= 0 or H <= 0:
        raise ValueError("Диаметр и высота должны быть положительными.")
    rows = _cells(table)
    if not rows:
        raise ValueError("Вставьте таблицу замеров (скопируйте блок чисел из Excel).")
    width = {len(r) for r in rows}
    if len(width) != 1:
        raise ValueError("В таблице строки разной длины. Скопируйте прямоугольный блок без заголовков.")
    grid = [[number(v, f"Таблица, строка {j + 1}, столбец {i + 1}") for i, v in enumerate(r)] for j, r in enumerate(rows)]
    if rows_are == "gens":                                 # транспонируем: строки = пояса
        grid = [list(col) for col in zip(*grid)]
    n_belts, n_gens = len(grid), len(grid[0])
    if n_gens < 3 or n_belts < 1:
        raise ValueError("Нужно минимум 3 образующие и 1 пояс.")
    if n_gens > MAX_GENS or n_belts > MAX_BELTS:
        raise ValueError("Слишком большая таблица.")
    th = parse_thicknesses(thicknesses)
    if len(th) != n_belts:
        raise ValueError(f"Толщин указано {len(th)}, а поясов в таблице {n_belts}. Количество должно совпадать.")
    step = H / n_belts
    heights = [0.0] + [round(step * (k + 1), 9) for k in range(n_belts)]
    heights[-1] = H
    angles = [360.0 * i / n_gens for i in range(n_gens)]
    values = [["0"] * n_gens] + [[repr(v) for v in row] for row in grid]
    return {"version": "1.0", "diameter": repr(D), "height": repr(H), "angles": angles, "heights": heights, "values": values,
            "bands": [[repr(step), repr(t)] for t in th], "defects": [], "contour_threshold": "0.1", "wall_includes_defects": True}


def _num(v):
    """Число из ячейки Excel: число или текст вида «-6,8» / «−6.8»; иначе None."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace("\u2212", "-").replace(",", ".").replace("\xa0", "").replace(" ", "")
    try:
        return float(t)
    except ValueError:
        return None


def _consecutive_from_one(vals) -> bool:
    return len(vals) >= 3 and all(abs(v - (i + 1)) < 1e-9 for i, v in enumerate(vals))


def table_from_xlsx(data: bytes) -> str:
    """Находит в первом листе Excel прямоугольный блок чисел (замеры) и возвращает его как текст с табуляцией.
    Заголовки, подписи и столбец/строку с номерами 1, 2, 3… отбрасывает."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ValueError("Не удалось открыть файл Excel. Загрузите файл .xlsx.") from e
    ws = wb.worksheets[0]
    rows = []
    for r in ws.iter_rows(values_only=True, max_row=600, max_col=250):
        nums = [_num(c) for c in r]
        rows.append(nums)
    wb.close()
    best, run = None, []
    def close(run):
        nonlocal best
        if len(run) >= 1 and len(run[0][1]) >= 2:
            score = len(run) * len(run[0][1])
            if best is None or score > best[0]:
                best = (score, run)
    prev = None
    for idx, nums in enumerate(rows):
        cols = tuple(i for i, v in enumerate(nums) if v is not None)
        if len(cols) >= 2 and cols == prev:
            run.append((idx, cols))
        else:
            close(run)
            run = [(idx, cols)] if len(cols) >= 2 else []
        prev = cols if len(cols) >= 2 else None
    close(run)
    if not best:
        raise ValueError("В файле не найдена таблица чисел. Нужен блок отклонений в мм (без пустых ячеек внутри).")
    run = best[1]
    cols = list(run[0][1])
    grid = [[rows[i][c] for c in cols] for i, _ in run]
    if len(grid[0]) > 2 and _consecutive_from_one([g[0] for g in grid]):          # столбец с номерами строк
        grid = [g[1:] for g in grid]
    if len(grid) > 2 and _consecutive_from_one(grid[0]):                           # строка с номерами столбцов
        grid = grid[1:]
    fmt = lambda v: str(int(v)) if float(v).is_integer() else repr(v)
    return "\n".join("\t".join(fmt(v) for v in g) for g in grid)


# ---------- геодезическая съёмка: книга Excel с листами «Т-001 полный», «Т-001 пустой» … ----------
def _survey_points(ws) -> dict[float, list[tuple[int, float, float]]]:
    """Блоки точек листа (разделены пустыми строками): уровень (мм) -> [(номер точки, отклонение мм, угол °)].
    Уровень блока - самое частое значение в столбце высоты (в съёмке встречаются округления 2980/2981). Точка 1 в конце блока - замыкание, пропускается."""
    from collections import Counter
    blocks, cur, lv = [], [], []

    def close():
        nonlocal cur, lv
        if len(cur) >= 3:
            blocks.append((Counter(round(x) for x in lv).most_common(1)[0][0], cur))
        cur, lv = [], []
    for r in ws.iter_rows(values_only=True, max_row=1500, max_col=10):
        if len(r) < 8:
            close(); continue
        no, dev, ang, level = r[1], _num(r[2]), _num(r[3]), _num(r[6])
        if isinstance(no, bool) or not isinstance(no, (int, float)) or dev is None or ang is None or level is None:
            close(); continue
        if int(no) == 1 and cur:                       # замыкание блока: точка 1 повторена
            continue
        cur.append((int(no), dev, ang))
        lv.append(level)
    close()
    out: dict[float, list] = {}
    for level, pts in blocks:
        out.setdefault(float(level), pts)
    return out


def survey_sheet_names(data: bytes) -> list[str]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    names = [ws.title for ws in wb.worksheets if len(_survey_points(ws)) >= 2]
    wb.close()
    return names


def survey_from_xlsx(data: bytes, sheet: str | None = None) -> dict:
    """Таблица замеров и размеры из листа съёмки. Возвращает {table (строки - образующие, столбцы - пояса), diameter_m, height_m, n_belts, n_gens, sheet}."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ValueError("Не удалось открыть файл Excel. Загрузите файл .xlsx.") from e
    cands = [ws for ws in wb.worksheets if len(_survey_points(ws)) >= 2]
    if not cands:
        wb.close()
        raise ValueError("survey:none")
    if sheet and sheet.strip():
        norm = lambda t: t.strip().lower().replace(" ", "").replace("\u2013", "-").replace("t", "\u0442")     # латинская «t» -> кириллическая «т»
        q = norm(sheet)
        pick = [ws for ws in cands if norm(ws.title) == q] or [ws for ws in cands if norm(ws.title).startswith(q)]
        if len(pick) > 1:                                  # «Т-001» без состояния: для модели берут «пустой» резервуар
            empty = [ws for ws in pick if "пуст" in ws.title.lower()]
            pick = empty if len(empty) == 1 else pick
        if len(pick) != 1:
            wb.close()
            raise ValueError(f"Лист «{sheet}» не найден или неоднозначен. Листы со съёмкой: {', '.join(w.title for w in cands)}.")
        ws = pick[0]
    elif len(cands) == 1:
        ws = cands[0]
    else:
        names = ", ".join(w.title for w in cands)
        wb.close()
        raise ValueError(f"В файле несколько листов со съёмкой: {names}. Впишите в поле «Лист» нужный (для модели обычно берут лист «пустой»).")
    pts = _survey_points(ws)
    r_mm = _num(ws["B1"].value)
    levels = sorted(pts)
    first = [p[2] for p in pts[levels[0]]]
    ng = len(first)
    for lv in levels:
        if len(pts[lv]) != ng or any(abs(a - b) > 0.5 for a, b in zip([p[2] for p in pts[lv]], first)):
            raise ValueError(f"На уровне {lv:g} мм число точек или углы не совпадают с первым уровнем.")
    steps = {round(abs(first[i + 1] - first[i]), 1) for i in range(ng - 1)}
    if len(steps) != 1 or abs(ng * steps.copy().pop() - 360) > 1.0:
        raise ValueError("Образующие должны быть расположены по окружности равномерно (одинаковый угловой шаг на 360°).")
    grid = [[pts[lv][g][1] for lv in levels] for g in range(ng)]           # [образующая][пояс]
    d_mm = (_num(ws["C1"].value) or (2 * r_mm if r_mm else None))
    if not d_mm:
        raise ValueError("В листе не найден диаметр (ячейки B1/C1).")
    table = "\n".join("\t".join(str(int(v)) if float(v).is_integer() else repr(v) for v in row) for row in grid)
    wb.close()
    return {"table": table, "diameter_m": d_mm / 1000, "height_m": levels[-1] / 1000, "n_belts": len(levels), "n_gens": ng, "sheet": ws.title}
